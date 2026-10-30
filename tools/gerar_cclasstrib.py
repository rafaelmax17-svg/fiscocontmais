# -*- coding: utf-8 -*-
"""Gera python/cclasstrib_dados.py a partir da planilha oficial do Informe Técnico 2025.002 (tabela cClassTrib do IBS/CBS).

Uso: python tools/gerar_cclasstrib.py "IT_2025.002_v.1.70_-_cClassTrib.xlsx" "v1.70"

Só a versão da planilha informada entra. Códigos com dFimVig preenchida ficam marcados com a data de fim; o módulo trata como
revogados a partir dela (não são apagados porque notas antigas ainda podem usá-los dentro da vigência).
"""
import sys
import os
import datetime
import pprint
import openpyxl


def _d(v):
    if isinstance(v, datetime.datetime):
        return v.strftime('%Y-%m-%d')
    return str(v)[:10] if v else ''


def _txt(v, n=900):
    v = ' '.join(str(v or '').split())
    return v if len(v) <= n else v[:n].rsplit(' ', 1)[0] + ' …'


def main(xlsx, versao):
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    ws = next(w for w in wb.worksheets if w.title.lower().startswith('cclass'))
    rows = list(ws.iter_rows(values_only=True))
    h = [str(x).strip() if x else '' for x in rows[0]]
    cods = {}
    for r in rows[1:]:
        d = dict(zip(h, r))
        c = str(d.get('cClassTrib') or '').strip()
        if not c:
            continue
        c = c.zfill(6)
        cods[c] = {
            'cst': str(d['CST-IBS/CBS']).zfill(3), 'cst_desc': _txt(d['Descrição CST-IBS/CBS'], 120),
            'nome': _txt(d['Nome cClassTrib'], 200), 'desc': _txt(d['Descrição cClassTrib'], 400),
            'artigo': _txt(d.get('LC 214/25'), 80), 'texto': _txt(d.get('LC Redação')),
            'tipo': _txt(d.get('Tipo de Alíquota'), 40),
            'pred_ibs': float(d.get('pRedIBS') or 0), 'pred_cbs': float(d.get('pRedCBS') or 0),
            'ini': _d(d.get('dIniVig')), 'fim': _d(d.get('dFimVig')),
            'nfe': int(d.get('indNFe') or 0), 'nfce': int(d.get('indNFCe') or 0), 'cte': int(d.get('indCTe') or 0),
        }
    atual = sorted({_d(dict(zip(h, r)).get('DataAtualização')) for r in rows[1:] if r[2]})
    ws2 = next(w for w in wb.worksheets if w.title.upper().startswith('CST'))
    csts = {}
    for r in list(ws2.iter_rows(values_only=True))[1:]:
        if r and r[0] is not None and str(r[0]).strip().isdigit():
            csts[str(r[0]).zfill(3)] = _txt(r[1], 120)
    dest = os.path.join(os.path.dirname(__file__), '..', 'python', 'cclasstrib_dados.py')
    with open(dest, 'w', encoding='utf-8') as fh:
        fh.write('# -*- coding: utf-8 -*-\n"""Tabela oficial cClassTrib do IBS/CBS (Informe Técnico 2025.002, ' + versao +
                 ', atualizada em ' + ', '.join(atual) + ').\nGerado por tools/gerar_cclasstrib.py — não editar à mão."""\n')
        fh.write('VERSAO = %r\nATUALIZACAO = %r\n' % (versao, atual[-1] if atual else ''))
        fh.write('CST = ' + pprint.pformat(csts, width=140) + '\n')
        fh.write('CCLASS = ' + pprint.pformat(cods, width=140) + '\n')
    print(len(cods), 'códigos,', sum(1 for v in cods.values() if v['fim']), 'com fim de vigência ·', len(csts), 'CSTs →', dest)


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else '')
