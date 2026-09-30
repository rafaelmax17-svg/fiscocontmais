# -*- coding: utf-8 -*-
"""
Gera python/efd_tabelas_natureza_dados.py a partir das tabelas oficiais da Receita
(Natureza da Receita da EFD-Contribuições: 4.3.10, 4.3.12, 4.3.13, 4.3.14, 4.3.15, 4.3.16).

Uso:  python tools/gerar_tabelas_natureza.py PASTA_COM_OS_DOCX
  - As tabelas saem do Portal SPED em .doc; abra no Word e "Salvar como .docx"
    (ou: soffice --headless --convert-to docx *.doc).
  - A tabela de cada arquivo é identificada pelo título ("Tabela 4.3.13 ..."), não pelo nome do arquivo.
  - Quando a Receita publicar versão nova, rode de novo com os arquivos novos e
    recompile: a versão/data de cada tabela aparece no painel.
"""
import docx, glob, os, re, sys, datetime

CSTS = {'4.3.10': ['02', '04'], '4.3.12': ['05'], '4.3.13': ['06'], '4.3.14': ['07'], '4.3.15': ['08'], '4.3.16': ['09']}


def limpa(s):
    return re.sub(r'\s+', ' ', (s or '').replace('\xa0', ' ')).strip()


def ler_tabela(caminho):
    d = docx.Document(caminho)
    titulo = limpa(d.paragraphs[0].text)
    m = re.search(r'4\.3\.\d+', titulo)
    if not m or m.group(0) not in CSTS:
        return None
    tab = m.group(0)
    partes = []
    v = re.search(r'Vers[ãa]o\s*([\d.]+)', titulo)
    a = re.search(r'Atualizada em\s*(\d{2}/\d{2}/\d{4})', titulo, re.I)
    if v: partes.append(f'versão {v.group(1)}')
    if a: partes.append(f'atualizada em {a.group(1)}')
    if not partes:
        outra = next((re.search(r'vers[ãa]o anterior \(([\d.]+)\)', limpa(p.text), re.I) for p in d.paragraphs
                      if re.search(r'vers[ãa]o anterior', p.text, re.I)), None)
        partes.append(f'posterior à versão {outra.group(1)}' if outra else 'versão não informada no arquivo')
    t = d.tables[0]
    cab = [limpa(c.text) for c in t.rows[0].cells]
    itens = {}
    for r in t.rows[1:]:
        c = [limpa(x.text) for x in r.cells]
        cod, desc = c[0], c[1] if len(c) > 1 else ''
        if not re.fullmatch(r'\d{3}', cod) or not desc:
            continue                                   # cabeçalhos de grupo, linhas vazias, subtítulos
        ncm = c[2] if len(c) > 2 else ''
        ini, fim = ('', '')
        if 'Início' in cab[-2]:
            ini, fim = c[-2], c[-1]
        elif len(c) > 6 and cab[-1] == '':             # 4.3.10: coluna extra vazia no fim
            ini, fim = c[5], c[6]
        if ini and not re.match(r'(\d{2}/)?\d{2}/\d{4}', ini):
            continue                                   # linha de título mesclada (texto repetido nas colunas)
        if re.fullmatch(r'\d00', cod) and not ini:
            continue                                   # título de grupo (100, 200...), não é código de natureza
        ncm = '' if ncm == '-' else ncm
        if cod in itens:                               # mesmo código em várias linhas (períodos de alíquota)
            if fim or not itens[cod]['t']:
                itens[cod]['t'] = fim
            continue
        itens[cod] = {'d': desc, 'n': ncm, 'i': ini, 't': fim}
    return tab, {'titulo': titulo, 'rotulo': ' · '.join(partes), 'csts': CSTS[tab], 'itens': itens}


def main(pasta):
    tabelas = {}
    for arq in sorted(glob.glob(os.path.join(pasta, '*.docx'))):
        r = ler_tabela(arq)
        if r:
            tabelas[r[0]] = r[1]
            print(f'{r[0]}  {len(r[1]["itens"]):>4} códigos  ({r[1]["rotulo"]})  <- {os.path.basename(arq)}')
    destino = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'python', 'efd_tabelas_natureza_dados.py')
    with open(destino, 'w', encoding='utf-8') as f:
        f.write('# -*- coding: utf-8 -*-\n# GERADO por tools/gerar_tabelas_natureza.py — não editar à mão.\n')
        f.write(f'# Fonte: tabelas oficiais da EFD-Contribuições (Receita Federal / Portal SPED). Gerado em {datetime.date.today():%d/%m/%Y}.\n')
        f.write('TABELAS = ' + repr(dict(sorted(tabelas.items()))) + '\n')
    print('gravado:', os.path.normpath(destino))


if __name__ == '__main__':
    main(sys.argv[1])
