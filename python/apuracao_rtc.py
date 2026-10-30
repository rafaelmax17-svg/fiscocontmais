# -*- coding: utf-8 -*-
"""Apuração Assistida IBS/CBS (Reforma Tributária) — a partir dos XMLs de NF-e e NFC-e.

Débitos: notas emitidas pela empresa (tpNF=1). Créditos: notas recebidas pela empresa e notas de entrada emitidas por ela
(tpNF=0). Os valores vêm do grupo IBSCBS de cada item (CST, cClassTrib, vBC, IBS UF, IBS município e CBS).

Base legal: LC 214/2025 — art. 46 (apuração assistida pela RFB e pelo Comitê Gestor), art. 47 (crédito nas aquisições,
exceto uso ou consumo pessoal), art. 57 (bens e serviços de uso ou consumo pessoal), arts. 343 e seguintes (2026, ano de
teste: CBS 0,9% e IBS 0,1%). Classificação conferida com a tabela oficial cClassTrib (Informe Técnico 2025.002, ver
cclasstrib_dados.py): código existente, vigente na data da nota, CST compatível, permitido no modelo, redução igual à da
tabela e valor = base × alíquota efetiva.

Regras que o sistema NÃO decide sozinho (ficam "Verificar"): crédito de item do art. 57 (bebida alcoólica, tabaco, joias,
armas, arte, veículos — pode ser revenda ou atividade, §3º), nota de fornecedor do regime normal sem o grupo IBS/CBS
(crédito estimado), nota de saída sem o grupo IBS/CBS. Nada nos XMLs é alterado.
"""
import os
import re
import json
import datetime
import xml.etree.ElementTree as ET
from collections import defaultdict, Counter

import cclasstrib_dados as TAB

NS = '{http://www.portalfiscal.inf.br/nfe}'
TOL = 0.011
CANCEL = '110111'

# Cenários (alíquotas nominais em %). 2026 = ano de teste (LC 214/2025, art. 343 e seguintes). Os demais são simulação:
# CBS 2027 = 11,3% calculado pelo TCU em 28/10/2026 e encaminhado ao Senado (fixação por resolução do Senado);
# IBS de referência estimado em 17% (editável na tela); 2029–2032 = 10%, 20%, 30% e 40% do IBS de referência.
CENARIOS = [
    {'ano': 2026, 'cbs': 0.9, 'ibs': 0.1, 'nota': 'Ano de teste (LC 214/2025, art. 343 e seguintes)'},
    {'ano': 2027, 'cbs': 11.3, 'ibs': 0.1, 'nota': 'CBS 11,3% calculada pelo TCU — depende de resolução do Senado'},
    {'ano': 2028, 'cbs': 11.3, 'ibs': 0.1, 'nota': 'Simulação'},
    {'ano': 2029, 'cbs': 11.3, 'ibs': 1.7, 'nota': 'IBS 10% da referência'},
    {'ano': 2030, 'cbs': 11.3, 'ibs': 3.4, 'nota': 'IBS 20% da referência'},
    {'ano': 2031, 'cbs': 11.3, 'ibs': 5.1, 'nota': 'IBS 30% da referência'},
    {'ano': 2032, 'cbs': 11.3, 'ibs': 6.8, 'nota': 'IBS 40% da referência'},
    {'ano': 2033, 'cbs': 11.3, 'ibs': 17.0, 'nota': 'IBS de referência estimado (17%)'},
]

# LC 214/2025, art. 57 — uso ou consumo pessoal (sem crédito, salvo uso preponderante na atividade, §3º)
ART57 = [(('2203', '2204', '2205', '2206', '2207', '2208'), 'bebida alcoólica'), (('24',), 'tabaco e derivados'),
         (('71',), 'joias, pedras e metais preciosos'), (('93',), 'armas e munições'), (('97',), 'obras de arte e antiguidades'),
         (('8703', '8711'), 'veículo')]
SIMPLES = {'1': 'Simples Nacional', '2': 'Simples Nacional (excesso de sublimite)', '4': 'MEI'}


def _f(s):
    try:
        return float(s) if s not in (None, '') else 0.0
    except ValueError:
        return 0.0


def _t(el, path):
    if el is None:
        return ''
    x = el.find(path)
    return (x.text or '').strip() if x is not None and x.text else ''


def _soma(el, path):
    return sum(_f(x.text) for x in el.findall(path)) if el is not None else 0.0


def _p(*partes):
    return '/'.join(NS + p for p in partes)


def art57(ncm):
    for pref, nome in ART57:
        if any(ncm.startswith(p) for p in pref):
            return nome
    return ''


# ------------------------------------------------------------------------------------------------------------ leitura
def ler_xml(texto):
    """Devolve ('nfe', dados) | ('cancel', chave) | (None, None)."""
    texto = re.sub(r'^\s*<\?xml[^>]*\?>', '', texto.lstrip('﻿'))
    try:
        raiz = ET.fromstring(texto)
    except ET.ParseError:
        return None, None
    ev = raiz.find('.//' + NS + 'infEvento')
    if ev is not None and raiz.find('.//' + NS + 'infNFe') is None:
        if _t(ev, NS + 'tpEvento') == CANCEL:
            ret = raiz.find('.//' + NS + 'retEvento/' + NS + 'infEvento')
            if ret is None or _t(ret, NS + 'cStat') in ('135', '136', '155', ''):
                return 'cancel', _t(ev, NS + 'chNFe')
        return None, None
    inf = raiz.find('.//' + NS + 'infNFe')
    if inf is None:
        return None, None
    chave = (inf.get('Id') or '')[3:]
    ide, emit, dest = inf.find(NS + 'ide'), inf.find(NS + 'emit'), inf.find(NS + 'dest')
    doc = {'chave': chave, 'mod': _t(ide, NS + 'mod'), 'serie': _t(ide, NS + 'serie'), 'num': _t(ide, NS + 'nNF'),
           'data': _t(ide, NS + 'dhEmi')[:10] or _t(ide, NS + 'dEmi')[:10], 'tpnf': _t(ide, NS + 'tpNF'),
           'fin': _t(ide, NS + 'finNFe'), 'emit': _t(emit, NS + 'CNPJ') or _t(emit, NS + 'CPF'),
           'emit_nome': _t(emit, NS + 'xNome'), 'crt': _t(emit, NS + 'CRT'),
           'dest': _t(dest, NS + 'CNPJ') or _t(dest, NS + 'CPF'), 'dest_nome': _t(dest, NS + 'xNome'),
           'vnf': _f(_t(inf, _p('total', 'ICMSTot', 'vNF'))), 'itens': []}
    for det in inf.findall(NS + 'det'):
        p, im = det.find(NS + 'prod'), det.find(NS + 'imposto')
        it = {'n': det.get('nItem'), 'cod': _t(p, NS + 'cProd'), 'desc': _t(p, NS + 'xProd'), 'ncm': _t(p, NS + 'NCM'),
              'cfop': _t(p, NS + 'CFOP'), 'vprod': _f(_t(p, NS + 'vProd')), 'vdesc': _f(_t(p, NS + 'vDesc')),
              'voutro': _f(_t(p, NS + 'vOutro')), 'vfrete': _f(_t(p, NS + 'vFrete')), 'vseg': _f(_t(p, NS + 'vSeg')),
              'vicms': _soma(im, NS + 'ICMS/*/' + NS + 'vICMS'),
              'vpc': _soma(im, NS + 'PIS/*/' + NS + 'vPIS') + _soma(im, NS + 'COFINS/*/' + NS + 'vCOFINS'),
              'grupo': False, 'cst': '', 'cclass': '', 'vbc': 0.0, 'vibsuf': 0.0, 'vibsmun': 0.0, 'vcbs': 0.0,
              'pibsuf': 0.0, 'pibsmun': 0.0, 'pcbs': 0.0, 'red': {}}
        ib = im.find(NS + 'IBSCBS') if im is not None else None
        if ib is not None:
            it['grupo'] = True
            it['cst'], it['cclass'] = _t(ib, NS + 'CST'), _t(ib, NS + 'cClassTrib')
            g = ib.find(NS + 'gIBSCBS')
            if g is not None:
                it['vbc'] = _f(_t(g, NS + 'vBC'))
                for chave_g, cam, campo_p, campo_v in (('ibsuf', 'gIBSUF', 'pIBSUF', 'vIBSUF'),
                                                     ('ibsmun', 'gIBSMun', 'pIBSMun', 'vIBSMun'),
                                                     ('cbs', 'gCBS', 'pCBS', 'vCBS')):
                    sub = g.find(NS + cam)
                    if sub is None:
                        continue
                    it['p' + chave_g] = _f(_t(sub, NS + campo_p))
                    it['v' + chave_g] = _f(_t(sub, NS + campo_v))
                    red = sub.find(NS + 'gRed')
                    if red is not None:
                        it['red'][chave_g] = (_f(_t(red, NS + 'pRedAliq')), _f(_t(red, NS + 'pAliqEfet')))
        doc['itens'].append(it)
    return 'nfe', doc


def ler_pastas(caminhos):
    from fiscal_core import coletar_xmls, _read_text
    docs, cancel, erros, lidos = {}, set(), 0, 0
    for c in coletar_xmls(caminhos):
        try:
            tipo, d = ler_xml(_read_text(c))
        except Exception:
            erros += 1
            continue
        lidos += 1
        if tipo == 'nfe' and d['chave']:
            docs[d['chave']] = d
        elif tipo == 'cancel' and d:
            cancel.add(d)
    return docs, cancel, lidos, erros


# ------------------------------------------------------------------------------------------------------------ conferência
def conferir_item(it, doc, papel, regime):
    """Lista de (situacao, texto) do item. situacao: Divergente | Verificar."""
    pts = []
    if not it['grupo']:
        return pts
    reg = TAB.CCLASS.get(it['cclass'])
    if not reg:
        return [('Divergente', f'cClassTrib {it["cclass"] or "(vazio)"} não existe na tabela oficial ({TAB.VERSAO})')]
    data = doc['data']
    if reg['ini'] and data and data < reg['ini']:
        pts.append(('Divergente', f'cClassTrib {it["cclass"]} só vale a partir de {_br(reg["ini"])} (nota de {_br(data)})'))
    if reg['fim'] and data and data >= reg['fim']:
        pts.append(('Divergente', f'cClassTrib {it["cclass"]} revogado: vigência encerrada em {_br(reg["fim"])}'))
    if it['cst'] != reg['cst']:
        pts.append(('Divergente', f'CST {it["cst"]} não combina com o cClassTrib {it["cclass"]} (tabela: CST {reg["cst"]})'))
    if doc['mod'] == '65' and not reg['nfce']:
        pts.append(('Divergente', f'cClassTrib {it["cclass"]} não pode ser usado em NFC-e (tabela oficial)'))
    if doc['mod'] == '55' and not reg['nfe']:
        pts.append(('Divergente', f'cClassTrib {it["cclass"]} não pode ser usado em NF-e (tabela oficial)'))
    for k, nome, pred in (('ibsuf', 'IBS UF', reg['pred_ibs']), ('ibsmun', 'IBS município', reg['pred_ibs']),
                          ('cbs', 'CBS', reg['pred_cbs'])):
        pr = it['red'].get(k)
        if pr and abs(pr[0] - pred) > 0.001:
            pts.append(('Divergente', f'Redução de {nome} informada {_n(pr[0])}% — a tabela manda {_n(pred)}%'))
        elif not pr and pred and it['p' + k] > 0 and it['v' + k] > 0:
            pts.append(('Divergente', f'{nome} sem a redução de {_n(pred)}% do cClassTrib {it["cclass"]}'))
        aliq = pr[1] if pr else it['p' + k]
        if it['vbc'] and abs(round(it['vbc'] * aliq / 100, 2) - it['v' + k]) > TOL:
            pts.append(('Divergente', f'{nome}: {_m(it["v" + k])} no XML, mas {_m(it["vbc"])} × {_n(aliq, 4)}% = '
                                      f'{_m(round(it["vbc"] * aliq / 100, 2))}'))
    if papel == 'debito' and regime and it['cclass'] != regime and reg['cst'] in ('000', '200'):
        r2 = TAB.CCLASS.get(regime, {})
        pts.append(('Divergente', f'Empresa marcada como "{r2.get("nome", regime)}" (cClassTrib {regime}, {r2.get("artigo", "")}) '
                                  f'mas a saída saiu com {it["cclass"]}'))
    return pts


def _br(d):
    return f'{d[8:10]}/{d[5:7]}/{d[:4]}' if d and len(d) >= 10 else d


def _n(v, casas=2):
    s = f'{v:,.{casas}f}'.replace(',', 'X').replace('.', ',').replace('X', '.')
    return s.rstrip('0').rstrip(',') if ',' in s else s


def _m(v):
    return 'R$ ' + f'{v:,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')


def base_estimada(it):
    # forma usada pela maioria dos emissores nos testes: valor do item + outras despesas + frete + seguro − desconto − ICMS − PIS − COFINS
    return round(it['vprod'] + it['voutro'] + it['vfrete'] + it['vseg'] - it['vdesc'] - it['vicms'] - it['vpc'], 2)


# ------------------------------------------------------------------------------------------------------------ SPED
def ler_sped(caminho):
    import difal_ro as dr
    texto, _ = dr._ler(caminho)
    cnpj, nome, ini, fim, notas = '', '', '', '', {}
    for l in texto.replace('\r\n', '\n').split('\n'):
        c = l.split('|')
        if len(c) < 3:
            continue
        if c[1] == '0000' and len(c) > 9:
            ini, fim, nome, cnpj = c[4], c[5], c[6], c[7]
        elif c[1] == 'C100' and len(c) > 9 and c[9]:
            notas[c[9]] = {'oper': c[2], 'sit': c[6], 'num': c[8], 'mod': c[5]}
        elif c[1] == '9999':
            break
    return {'cnpj': cnpj, 'nome': nome, 'ini': ini, 'fim': fim, 'notas': notas}


# ------------------------------------------------------------------------------------------------------------ análise
def analisar(caminhos, cnpj='', regime='', sped=''):
    docs, cancel, lidos, erros = ler_pastas(caminhos)
    sp = ler_sped(sped) if sped else None
    if sp and not cnpj:
        cnpj = sp['cnpj']
    cnpj = re.sub(r'\D', '', cnpj or '')
    if not cnpj:
        cont = Counter()
        for d in docs.values():
            cont[d['emit']] += 2 if d['tpnf'] == '1' else 1
            cont[d['dest']] += 1
        cont.pop('', None)
        cnpj = cont.most_common(1)[0][0] if cont else ''
    nome = (sp or {}).get('nome', '')
    if regime and regime not in TAB.CCLASS:
        regime = ''

    tot = defaultdict(float)
    sim = defaultdict(float)        # bases efetivas para a simulação (base × (1 − redução))
    por_cfop = defaultdict(lambda: defaultdict(float))
    itens_out, docs_out, pontos = [], [], []
    produtos = {}
    meses = Counter()
    sem_grupo_simples = []
    ignorados = Counter()
    cont_sit = Counter()

    for d in sorted(docs.values(), key=lambda x: (x['data'], x['num'])):
        if d['chave'] in cancel:
            ignorados['canceladas (evento no lote)'] += 1
            continue
        if d['emit'] == cnpj and d['tpnf'] == '1':
            papel = 'debito'
        elif d['dest'] == cnpj and d['emit'] != cnpj:
            papel = 'credito' if d['tpnf'] == '1' else None
        elif d['emit'] == cnpj and d['tpnf'] == '0':
            papel = 'credito'
        else:
            papel = None
        if not papel:
            ignorados['de outra empresa ou sem relação com o CNPJ'] += 1
            continue
        if d['dest'] == cnpj and d['tpnf'] == '0' and d['emit'] != cnpj:
            ignorados['entrada emitida pelo fornecedor (tpNF=0)'] += 1
            continue
        meses[d['data'][:7]] += 1
        dsum = defaultdict(float)
        d_sit = Counter()
        simples = papel == 'credito' and d['emit'] != cnpj and d['crt'] in SIMPLES
        tem_grupo = any(i['grupo'] for i in d['itens'])
        if simples and not tem_grupo:
            sem_grupo_simples.append({'chave': d['chave'], 'num': d['num'], 'data': d['data'], 'emit': d['emit'],
                                      'nome': d['emit_nome'], 'regime': SIMPLES[d['crt']], 'vnf': d['vnf']})
        for it in d['itens']:
            pts = conferir_item(it, d, papel, regime)
            origem = 'XML'
            vibs, vcbs, vbc = it['vibsuf'] + it['vibsmun'], it['vcbs'], it['vbc']
            reg = TAB.CCLASS.get(it['cclass'], {})
            f_ibs = 1 - reg.get('pred_ibs', 0) / 100
            f_cbs = 1 - reg.get('pred_cbs', 0) / 100
            tributa = reg.get('cst') not in ('410', '400', '620', '550', '510') if reg else True
            if not it['grupo']:
                if simples:
                    origem = 'Simples (sem crédito)'
                    vibs = vcbs = vbc = 0.0
                    tributa = False
                elif it['ncm'].startswith(('2710', '2711', '2207', '3826')):
                    origem = 'sem valor'
                    vibs = vcbs = vbc = 0.0
                    tributa = False
                    pts.append(('Verificar', 'Combustível sem o grupo IBS/CBS no XML — tributação monofásica (LC 214/2025, '
                                             'art. 172 e seguintes): o sistema não estima valor'))
                else:
                    origem = 'estimado'
                    vbc = base_estimada(it)
                    vibs, vcbs = round(vbc * 0.1 / 100, 2), round(vbc * 0.9 / 100, 2)
                    f_ibs = f_cbs = 1.0
                    pts.append(('Verificar', ('Saída sem o grupo IBS/CBS no XML — valor estimado com tributação integral de 2026'
                                              if papel == 'debito' else
                                              'Fornecedor do regime normal sem o grupo IBS/CBS — crédito estimado com tributação '
                                              'integral de 2026 (se o produto tiver redução, o crédito é menor)')))
            a57 = art57(it['ncm']) if papel == 'credito' else ''
            confirmado = True
            if a57 and (vibs or vcbs):
                confirmado = False
                pts.append(('Verificar', f'{a57.capitalize()}: uso ou consumo pessoal não dá crédito (LC 214/2025, art. 57); '
                                         'se for revenda ou usado na atividade, o crédito vale (§3º)'))
            sit = 'Divergente' if any(s == 'Divergente' for s, _ in pts) else ('Verificar' if pts else
                                                                               ('Sem IBS/CBS' if not it['grupo'] else 'Confere'))
            cont_sit[sit] += 1
            d_sit[sit] += 1
            sinal = 1 if papel == 'debito' else -1
            chave_tot = papel if confirmado and origem == 'XML' else (
                papel + '_est' if origem == 'estimado' else (papel + '_57' if not confirmado else papel + '_simples'))
            tot[chave_tot + '_ibs'] += vibs
            tot[chave_tot + '_cbs'] += vcbs
            tot[chave_tot + '_bc'] += vbc
            tot[chave_tot + '_ibsuf'] += it['vibsuf']
            tot[chave_tot + '_ibsmun'] += it['vibsmun']
            if tributa and confirmado and (origem == 'XML' or (origem == 'estimado' and papel == 'debito')):
                sim[papel + '_ibs'] += vbc * f_ibs
                sim[papel + '_cbs'] += vbc * f_cbs
            pc = por_cfop[(papel, it['cfop'])]
            pc['bc'] += vbc
            pc['ibs'] += vibs
            pc['cbs'] += vcbs
            pc['itens'] += 1
            dsum['bc'] += vbc
            dsum['ibs'] += vibs
            dsum['cbs'] += vcbs
            chp = (papel, it['cod'] or it['desc'], it['ncm'], it['cst'], it['cclass'])
            pr = produtos.get(chp)
            if not pr:
                pr = produtos[chp] = {'papel': papel, 'cod': it['cod'], 'desc': it['desc'], 'ncm': it['ncm'],
                                      'cst': it['cst'], 'cclass': it['cclass'], 'nome': reg.get('nome', ''),
                                      'artigo': reg.get('artigo', ''), 'texto': reg.get('texto', ''),
                                      'itens': 0, 'bc': 0.0, 'ibs': 0.0, 'cbs': 0.0, 'sit': Counter(), 'motivos': Counter(),
                                      'notas': set(), 'origem': Counter()}
            pr['itens'] += 1
            pr['bc'] += vbc
            pr['ibs'] += vibs
            pr['cbs'] += vcbs
            pr['sit'][sit] += 1
            pr['origem'][origem] += 1
            pr['notas'].add(d['num'])
            for s, txt in pts:
                pr['motivos'][(s, txt if 'no XML, mas' not in txt else txt.split(':')[0] + ': valor diferente de base × alíquota')] += 1
            for s, txt in pts:
                pontos.append({'sit': s, 'papel': papel, 'num': d['num'], 'mod': d['mod'], 'data': d['data'],
                               'chave': d['chave'], 'item': it['n'], 'cod': it['cod'], 'desc': it['desc'], 'ncm': it['ncm'],
                               'cst': it['cst'], 'cclass': it['cclass'], 'texto': txt})
            if len(itens_out) < 60000:
                itens_out.append({'papel': papel, 'num': d['num'], 'mod': d['mod'], 'data': d['data'], 'chave': d['chave'],
                                  'part': d['dest_nome'] if papel == 'debito' else d['emit_nome'],
                                  'item': it['n'], 'cod': it['cod'], 'desc': it['desc'], 'ncm': it['ncm'], 'cfop': it['cfop'],
                                  'cst': it['cst'], 'cclass': it['cclass'], 'bc': round(vbc, 2),
                                  'ibsuf': round(it['vibsuf'], 2), 'ibsmun': round(it['vibsmun'], 2), 'ibs': round(vibs, 2),
                                  'cbs': round(vcbs, 2), 'origem': origem, 'sit': sit,
                                  'motivo': ' · '.join(t for _, t in pts)})
        docs_out.append({'papel': papel, 'mod': d['mod'], 'num': d['num'], 'serie': d['serie'], 'data': d['data'],
                         'chave': d['chave'], 'part': d['dest_nome'] if papel == 'debito' else d['emit_nome'],
                         'part_doc': d['dest'] if papel == 'debito' else d['emit'], 'crt': d['crt'], 'vnf': d['vnf'],
                         'bc': round(dsum['bc'], 2), 'ibs': round(dsum['ibs'], 2), 'cbs': round(dsum['cbs'], 2),
                         'sit': 'Divergente' if d_sit['Divergente'] else ('Verificar' if d_sit['Verificar'] else
                                                                          ('Sem IBS/CBS' if d_sit['Sem IBS/CBS'] else 'Confere'))})

    r2 = lambda v: round(v, 2)
    deb_ibs, deb_cbs = r2(tot['debito_ibs'] + tot['debito_est_ibs']), r2(tot['debito_cbs'] + tot['debito_est_cbs'])
    cred_ibs, cred_cbs = r2(tot['credito_ibs']), r2(tot['credito_cbs'])
    est_ibs, est_cbs = r2(tot['credito_est_ibs']), r2(tot['credito_est_cbs'])
    a57_ibs, a57_cbs = r2(tot['credito_57_ibs']), r2(tot['credito_57_cbs'])
    painel = {
        'debito': {'ibs': deb_ibs, 'cbs': deb_cbs, 'total': r2(deb_ibs + deb_cbs),
                   'ibsuf': r2(tot['debito_ibsuf']), 'ibsmun': r2(tot['debito_ibsmun']),
                   'bc': r2(tot['debito_bc'] + tot['debito_est_bc']), 'estimado': r2(tot['debito_est_ibs'] + tot['debito_est_cbs'])},
        'credito': {'ibs': cred_ibs, 'cbs': cred_cbs, 'total': r2(cred_ibs + cred_cbs),
                    'ibsuf': r2(tot['credito_ibsuf']), 'ibsmun': r2(tot['credito_ibsmun']), 'bc': r2(tot['credito_bc'])},
        'estimado': {'ibs': est_ibs, 'cbs': est_cbs, 'total': r2(est_ibs + est_cbs), 'bc': r2(tot['credito_est_bc'])},
        'art57': {'ibs': a57_ibs, 'cbs': a57_cbs, 'total': r2(a57_ibs + a57_cbs), 'bc': r2(tot['credito_57_bc'])},
        'simples': {'bc': r2(tot['credito_simples_bc']), 'notas': len(sem_grupo_simples)},
    }
    painel['saldo'] = {'ibs': r2(deb_ibs - cred_ibs), 'cbs': r2(deb_cbs - cred_cbs),
                       'total': r2(deb_ibs + deb_cbs - cred_ibs - cred_cbs)}
    painel['saldo_com_ajustes'] = r2(painel['saldo']['total'] - painel['estimado']['total'] - painel['art57']['total'])

    xs = None
    if sp:
        chaves_xml = {d['chave'] for d in docs.values() if (d['emit'] == cnpj or d['dest'] == cnpj)}
        canc_sped = {k for k, v in sp['notas'].items() if v['sit'] in ('02', '03')}
        sem_sped = sorted(k for k in chaves_xml - set(sp['notas']) if k not in cancel)
        sem_xml = sorted(k for k, v in sp['notas'].items() if k not in chaves_xml and v['sit'] not in ('02', '03', '04', '05')
                         and v['mod'] in ('55', '65'))
        canc_pasta = sorted(k for k in canc_sped if k in chaves_xml and k not in cancel)
        dd = lambda k: {'chave': k, 'num': (docs.get(k) or {}).get('num') or sp['notas'].get(k, {}).get('num', ''),
                        'data': (docs.get(k) or {}).get('data', ''), 'vnf': (docs.get(k) or {}).get('vnf', 0.0),
                        'tipo': ('saída' if (sp['notas'].get(k, {}).get('oper') == '1' or (docs.get(k) or {}).get('emit') == cnpj)
                                 else 'entrada')}
        xs = {'periodo': f"{_br8(sp['ini'])} a {_br8(sp['fim'])}", 'empresa': sp['nome'],
              'sem_sped': [dd(k) for k in sem_sped], 'sem_xml': [dd(k) for k in sem_xml],
              'canc_pasta': [dd(k) for k in canc_pasta], 'cnpj_difere': bool(sp['cnpj'] and sp['cnpj'] != cnpj)}

    if not nome:
        nome = next((d['emit_nome'] for d in docs.values() if d['emit'] == cnpj), '') or \
               next((d['dest_nome'] for d in docs.values() if d['dest'] == cnpj), '')
    prods = []
    for p in produtos.values():
        sit = 'Divergente' if p['sit']['Divergente'] else ('Verificar' if p['sit']['Verificar'] else
                                                           ('Sem IBS/CBS' if p['sit']['Sem IBS/CBS'] else 'Confere'))
        prods.append({'papel': p['papel'], 'cod': p['cod'], 'desc': p['desc'], 'ncm': p['ncm'], 'cst': p['cst'],
                      'cclass': p['cclass'], 'nome': p['nome'], 'artigo': p['artigo'], 'texto': p['texto'],
                      'itens': p['itens'], 'notas': len(p['notas']), 'bc': r2(p['bc']), 'ibs': r2(p['ibs']),
                      'cbs': r2(p['cbs']), 'sit': sit, 'origem': ', '.join(p['origem']),
                      'motivos': [{'sit': s, 'texto': t, 'qtd': q} for (s, t), q in p['motivos'].most_common()]})
    ordem = {'Divergente': 0, 'Verificar': 1, 'Sem IBS/CBS': 2, 'Confere': 3}
    prods.sort(key=lambda x: (ordem[x['sit']], x['papel'], -(x['ibs'] + x['cbs'])))
    cls = Counter((i['cst'], i['cclass']) for i in itens_out if i['cclass'])
    classes = [{'cst': c, 'cclass': k, 'itens': q, 'nome': TAB.CCLASS.get(k, {}).get('nome', 'não existe na tabela'),
                'red': TAB.CCLASS.get(k, {}).get('pred_cbs', 0), 'artigo': TAB.CCLASS.get(k, {}).get('artigo', '')}
               for (c, k), q in cls.most_common()]
    return {
        'cnpj': cnpj, 'empresa': nome, 'regime': regime, 'regime_nome': TAB.CCLASS.get(regime, {}).get('nome', '') if regime else '',
        'meses': sorted(meses), 'xmls_lidos': lidos, 'erros': erros, 'canceladas': len(cancel), 'ignorados': dict(ignorados),
        'notas': {'debito': sum(1 for d in docs_out if d['papel'] == 'debito'),
                  'credito': sum(1 for d in docs_out if d['papel'] == 'credito')},
        'itens_total': sum(cont_sit.values()), 'situacao': dict(cont_sit),
        'painel': painel, 'simulacao': {'cenarios': CENARIOS, 'bases': {k: r2(v) for k, v in sim.items()}},
        'por_cfop': sorted([{'papel': k[0], 'cfop': k[1], 'bc': r2(v['bc']), 'ibs': r2(v['ibs']), 'cbs': r2(v['cbs']),
                             'itens': int(v['itens'])} for k, v in por_cfop.items()], key=lambda x: (x['papel'] != 'debito', -x['bc'])),
        'documentos': docs_out, 'produtos': prods, 'classes': classes, 'itens': itens_out,
        'pontos': pontos, 'simples_sem_grupo': sem_grupo_simples, 'xml_sped': xs,
        'tabela': {'versao': TAB.VERSAO, 'atualizacao': TAB.ATUALIZACAO, 'codigos': len(TAB.CCLASS),
                   'revogados': sum(1 for v in TAB.CCLASS.values() if v['fim'])},
        'regimes_especificos': [{'cclass': k, 'nome': v['nome'], 'artigo': v['artigo']} for k, v in sorted(TAB.CCLASS.items())
                                if v['cst'] == '200' and (v['nfe'] or v['nfce']) and not v['fim'] and v['pred_cbs'] < 100
                                and any(w in v['nome'].lower() for w in ('bares', 'hotel', 'turismo', 'profiss', 'imóve', 'transporte'))],
    }


def _br8(d):
    return f'{d[:2]}/{d[2:4]}/{d[4:]}' if d and len(d) == 8 else d


# ------------------------------------------------------------------------------------------------------------ Excel
def gerar_xlsx(res, destino):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    TINTA, LAR, VERDE, PAPEL = '1C1917', 'E2552D', '0F766E', 'F4F1EA'
    cab_f, cab_font = PatternFill('solid', fgColor='EFEAE0'), Font(bold=True, color='57534E', size=10)
    linha = Border(bottom=Side(style='thin', color='E7E2D8'))
    cores = {'Divergente': ('FECACA', '991B1B'), 'Verificar': ('FDE68A', '78350F'), 'Confere': ('CCFBF1', '0F766E'),
             'Sem IBS/CBS': ('E7E5E4', '44403C')}
    din = '#,##0.00'

    def tabela(ws, r0, cab, linhas, larg, moedas=()):
        for j, c in enumerate(cab, 1):
            x = ws.cell(r0, j, c)
            x.fill, x.font, x.alignment = cab_f, cab_font, Alignment(vertical='center', wrap_text=True)
        for i, l in enumerate(linhas, r0 + 1):
            for j, v in enumerate(l, 1):
                x = ws.cell(i, j, v)
                x.border = linha
                if j - 1 in moedas:
                    x.number_format = din
                if isinstance(v, str) and v in cores:
                    x.fill = PatternFill('solid', fgColor=cores[v][0])
                    x.font = Font(bold=True, color=cores[v][1], size=10)
        for j, w in enumerate(larg, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
        ws.freeze_panes = ws.cell(r0 + 1, 1)
        if linhas:
            ws.auto_filter.ref = f'A{r0}:{get_column_letter(len(cab))}{r0 + len(linhas)}'

    p = res['painel']
    ws = wb.active
    ws.title = 'Painel'
    ws.sheet_view.showGridLines = False
    for col, w in zip('ABCDEFG', (3, 34, 18, 18, 18, 3, 3)):
        ws.column_dimensions[col].width = w
    for r in range(1, 40):
        for c in range(1, 8):
            ws.cell(r, c).fill = PatternFill('solid', fgColor=PAPEL)
    ws['B2'] = 'Apuração Assistida IBS / CBS'
    ws['B2'].font = Font(bold=True, size=20, color=TINTA)
    ws['B3'] = 'REFORMA TRIBUTÁRIA'
    ws['B3'].font = Font(bold=True, size=9, color='FFFFFF')
    ws['B3'].fill = PatternFill('solid', fgColor=LAR)
    ws['B4'] = f"{res['empresa']} · CNPJ {res['cnpj']} · competência {', '.join(_mes(m) for m in res['meses'])}"
    ws['B4'].font = Font(size=10, color='57534E')
    lin = [('Débitos (saídas)', p['debito']['cbs'], p['debito']['ibs'], p['debito']['total'], LAR),
           ('Créditos destacados (entradas)', -p['credito']['cbs'], -p['credito']['ibs'], -p['credito']['total'], VERDE),
           ('Saldo apurado', p['saldo']['cbs'], p['saldo']['ibs'], p['saldo']['total'], TINTA)]
    for j, h in enumerate(('', 'CBS', 'IBS', 'Total'), 2):
        c = ws.cell(6, j, h)
        c.font = Font(bold=True, color='78716C', size=9)
    for i, (rot, cbs, ibs, tt, cor) in enumerate(lin, 7):
        ws.cell(i, 2, rot).font = Font(bold=True, color=cor, size=12)
        for j, v in enumerate((cbs, ibs, tt), 3):
            c = ws.cell(i, j, v)
            c.number_format = din
            c.font = Font(bold=True, color=cor, size=12)
            c.border = Border(bottom=Side(style='thin', color='D6D3D1'))
    ws['B11'] = 'Ajustes que dependem de decisão'
    ws['B11'].font = Font(bold=True, size=11, color=TINTA)
    aj = [('Crédito estimado (fornecedor do regime normal sem IBS/CBS no XML)', p['estimado']['total']),
          ('Crédito a confirmar — art. 57 (uso ou consumo pessoal × atividade)', p['art57']['total']),
          ('Saldo se os dois forem confirmados', p['saldo_com_ajustes'])]
    for i, (rot, v) in enumerate(aj, 12):
        ws.cell(i, 2, rot).font = Font(size=10, color='44403C')
        c = ws.cell(i, 5, v)
        c.number_format = din
        c.font = Font(bold=True, size=10)
    ws['B16'] = 'Situação dos itens'
    ws['B16'].font = Font(bold=True, size=11, color=TINTA)
    for i, s in enumerate(('Confere', 'Divergente', 'Verificar', 'Sem IBS/CBS'), 17):
        c = ws.cell(i, 2, s)
        c.fill = PatternFill('solid', fgColor=cores[s][0])
        c.font = Font(bold=True, color=cores[s][1])
        ws.cell(i, 3, res['situacao'].get(s, 0))
    ws['B22'] = ('Base legal: LC 214/2025 — art. 46 (apuração assistida), art. 47 (créditos), art. 57 (uso ou consumo pessoal), '
                 f"art. 343 e seguintes (ano de teste). Tabela cClassTrib oficial {res['tabela']['versao']} "
                 f"(atualizada em {_br(res['tabela']['atualizacao'])}).")
    ws['B22'].font = Font(size=9, color='78716C', italic=True)

    w = wb.create_sheet('Por CFOP')
    tabela(w, 1, ['Tipo', 'CFOP', 'Itens', 'Base IBS/CBS', 'IBS', 'CBS'],
           [['Débito' if x['papel'] == 'debito' else 'Crédito', x['cfop'], x['itens'], x['bc'], x['ibs'], x['cbs']]
            for x in res['por_cfop']], [10, 8, 8, 16, 14, 14], (3, 4, 5))
    w = wb.create_sheet('Por documento')
    tabela(w, 1, ['Tipo', 'Modelo', 'Número', 'Data', 'Participante', 'CNPJ/CPF', 'Valor da nota', 'Base', 'IBS', 'CBS', 'Situação', 'Chave'],
           [['Débito' if x['papel'] == 'debito' else 'Crédito', x['mod'], x['num'], _br(x['data']), x['part'], x['part_doc'],
             x['vnf'], x['bc'], x['ibs'], x['cbs'], x['sit'], x['chave']] for x in res['documentos']],
           [9, 8, 10, 11, 34, 17, 14, 14, 11, 11, 13, 47], (6, 7, 8, 9))
    w = wb.create_sheet('Classificação')
    tabela(w, 1, ['Tipo', 'Código', 'Produto', 'NCM', 'CST', 'cClassTrib', 'Classificação na tabela oficial', 'Artigo',
                  'Itens', 'Base', 'IBS', 'CBS', 'Situação', 'Motivo', 'Fundamento (texto da lei)'],
           [['Débito' if x['papel'] == 'debito' else 'Crédito', x['cod'], x['desc'], x['ncm'], x['cst'], x['cclass'], x['nome'],
             x['artigo'], x['itens'], x['bc'], x['ibs'], x['cbs'], x['sit'],
             ' · '.join(f"{m['texto']} ({m['qtd']})" for m in x['motivos']), x['texto']] for x in res['produtos']],
           [9, 10, 34, 11, 6, 10, 34, 10, 7, 13, 11, 11, 13, 60, 70], (9, 10, 11))
    w = wb.create_sheet('Pontos de atenção')
    tabela(w, 1, ['Situação', 'Tipo', 'Nota', 'Data', 'Item', 'Produto', 'NCM', 'CST', 'cClassTrib', 'O que foi encontrado', 'Chave'],
           [[x['sit'], 'Débito' if x['papel'] == 'debito' else 'Crédito', x['num'], _br(x['data']), x['item'], x['desc'], x['ncm'],
             x['cst'], x['cclass'], x['texto'], x['chave']] for x in res['pontos'][:100000]],
           [12, 9, 10, 11, 6, 34, 11, 6, 10, 80, 47])
    w = wb.create_sheet('Itens')
    tabela(w, 1, ['Tipo', 'Nota', 'Data', 'Item', 'Código', 'Produto', 'NCM', 'CFOP', 'CST', 'cClassTrib', 'Base', 'IBS UF',
                  'IBS mun.', 'CBS', 'Origem', 'Situação', 'Motivo'],
           [['Débito' if x['papel'] == 'debito' else 'Crédito', x['num'], _br(x['data']), x['item'], x['cod'], x['desc'], x['ncm'],
             x['cfop'], x['cst'], x['cclass'], x['bc'], x['ibsuf'], x['ibsmun'], x['cbs'], x['origem'], x['sit'], x['motivo']]
            for x in res['itens']],
           [9, 10, 11, 6, 10, 34, 11, 7, 6, 10, 13, 11, 11, 11, 14, 13, 60], (10, 11, 12, 13))
    if res['simples_sem_grupo']:
        w = wb.create_sheet('Simples sem crédito')
        tabela(w, 1, ['Nota', 'Data', 'Fornecedor', 'CNPJ', 'Regime', 'Valor'],
               [[x['num'], _br(x['data']), x['nome'], x['emit'], x['regime'], x['vnf']] for x in res['simples_sem_grupo']],
               [10, 11, 40, 17, 26, 14], (5,))
    w = wb.create_sheet('Simulação até 2033')
    b = res['simulacao']['bases']
    lin = []
    for c in res['simulacao']['cenarios']:
        deb = b.get('debito_cbs', 0) * c['cbs'] / 100 + b.get('debito_ibs', 0) * c['ibs'] / 100
        cred = b.get('credito_cbs', 0) * c['cbs'] / 100 + b.get('credito_ibs', 0) * c['ibs'] / 100
        lin.append([c['ano'], c['cbs'] / 100, c['ibs'] / 100, round(deb, 2), round(cred, 2), round(deb - cred, 2), c['nota']])
    tabela(w, 1, ['Ano', 'CBS', 'IBS', 'Débitos', 'Créditos', 'Saldo', 'Observação'], lin, [8, 9, 9, 15, 15, 15, 60], (3, 4, 5))
    for r in range(2, len(lin) + 2):
        w.cell(r, 2).number_format = w.cell(r, 3).number_format = '0.0%'
    w.cell(len(lin) + 3, 1, 'Simulação com a mesma movimentação do período. Não são valores oficiais; a CBS de 2027 depende de '
                            'resolução do Senado e o IBS de referência ainda não foi fixado.').font = Font(italic=True, size=9, color='78716C')
    if res.get('xml_sped'):
        x = res['xml_sped']
        w = wb.create_sheet('XML × SPED')
        rows = [['XML sem escrituração no SPED', k['tipo'], k['num'], _br(k['data']), k['vnf'], k['chave']] for k in x['sem_sped']]
        rows += [['Escriturada no SPED sem XML no lote', k['tipo'], k['num'], _br(k['data']), k['vnf'], k['chave']] for k in x['sem_xml']]
        rows += [['Cancelada no SPED e autorizada no lote', k['tipo'], k['num'], _br(k['data']), k['vnf'], k['chave']] for k in x['canc_pasta']]
        tabela(w, 1, ['Ocorrência', 'Tipo', 'Nota', 'Data', 'Valor', 'Chave'], rows, [38, 9, 10, 11, 14, 47], (4,))
    tmp = destino + '.parcial'
    wb.save(tmp)
    os.replace(tmp, destino)


def _mes(m):
    return f'{m[5:7]}/{m[:4]}' if len(m) == 7 else m


# ------------------------------------------------------------------------------------------------------------ HTML
def gerar_html(res):
    dados = {k: res[k] for k in ('cnpj', 'empresa', 'regime', 'regime_nome', 'meses', 'notas', 'itens_total', 'situacao',
                                 'painel', 'simulacao', 'por_cfop', 'classes', 'simples_sem_grupo', 'xml_sped', 'tabela',
                                 'ignorados', 'canceladas', 'xmls_lidos')}
    dados['produtos'] = res['produtos'][:5000]
    dados['documentos'] = res['documentos'][:8000]
    dados['pontos'] = res['pontos'][:8000]
    dados['n_pontos'] = len(res['pontos'])
    js = json.dumps(dados, ensure_ascii=False).replace('</', '<\\/')
    return HTML.replace('/*DADOS*/', js)


HTML = r'''<!doctype html><html lang="pt-BR"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Apuração Assistida IBS/CBS</title><style>
:root{--papel:#f4f1ea;--card:#fff;--linha:#e7e2d8;--tinta:#1c1917;--suave:#57534e;--fraco:#78716c;--lar:#e2552d;--verde:#0f766e;--ouro:#fbbf24}
*{box-sizing:border-box}body{margin:0;font-family:"Segoe UI",Arial,sans-serif;background:var(--papel);color:var(--tinta)}
.wrap{padding:24px 28px;max-width:1500px;margin:0 auto}
.hd{display:flex;justify-content:space-between;align-items:flex-end;gap:16px;border-bottom:2px solid var(--tinta);padding-bottom:12px;flex-wrap:wrap}
.hd h1{margin:0;font-size:26px;font-weight:800;letter-spacing:-.5px}
.rt{font-size:11px;font-weight:800;letter-spacing:1.2px;padding:4px 10px;border-radius:4px;background:var(--lar);color:#fff;margin-left:10px;vertical-align:middle}
.hd .m{font-size:12.5px;color:var(--suave);text-align:right;line-height:1.5}
.tabs{display:flex;gap:22px;margin:14px 0 18px;font-size:13.5px;font-weight:600;color:var(--fraco);flex-wrap:wrap}
.tabs span{cursor:pointer;padding-bottom:4px;border-bottom:3px solid transparent}.tabs span.on{color:var(--tinta);border-color:var(--lar)}
.pg{display:none;animation:sobe .45s ease}.pg.on{display:block}
@keyframes sobe{from{opacity:0;transform:translateY(8px)}to{opacity:1;transform:none}}
.bento{display:grid;grid-template-columns:1.25fr 1fr 1fr;gap:14px}
.c{background:var(--card);border:1px solid var(--linha);border-radius:16px;padding:16px 18px}
.c h3{margin:0 0 4px;font-size:11.5px;text-transform:uppercase;letter-spacing:1px;color:var(--fraco)}
.big{font-size:34px;font-weight:800;letter-spacing:-1px}.lbl{font-size:12px;color:var(--suave)}
.dark{background:var(--tinta);color:var(--papel);border:0}.dark h3{color:#a8a29e}
.chip{display:inline-block;font-size:11px;padding:3px 9px;border-radius:20px;margin:4px 4px 0 0;background:#44403c;color:#fde68a}
.legend{display:flex;gap:14px;font-size:11.5px;color:var(--suave);margin-top:6px;flex-wrap:wrap}.legend i{display:inline-block;width:10px;height:10px;border-radius:2px;margin-right:5px;vertical-align:-1px}
.aj{display:flex;justify-content:space-between;font-size:12.5px;padding:7px 0;border-top:1px dashed var(--linha)}.aj b{white-space:nowrap}
table{width:100%;border-collapse:separate;border-spacing:0;background:#fff;border:1px solid var(--linha);border-radius:14px;overflow:hidden;font-size:12.5px}
th{text-align:left;padding:9px 10px;background:#efeae0;color:var(--suave);font-weight:700;font-size:11px;text-transform:uppercase;letter-spacing:.4px;position:sticky;top:0}
td{padding:8px 10px;border-top:1px solid #efeae0;vertical-align:top}td.n,th.n{text-align:right;white-space:nowrap}
.code{font-family:Consolas,monospace;font-weight:700}.sm{font-size:11px;color:var(--fraco)}
.tg{font-size:10.5px;padding:2px 8px;border-radius:20px;font-weight:700;white-space:nowrap}
.t-Confere{background:#ccfbf1;color:#0f766e}.t-Verificar{background:#fde68a;color:#78350f}.t-Divergente{background:#fecaca;color:#991b1b}.t-Sem{background:#e7e5e4;color:#44403c}
.bar{display:flex;gap:10px;flex-wrap:wrap;margin-bottom:12px;align-items:center}
.st{background:#fff;border:1px solid var(--linha);border-radius:12px;padding:8px 14px;font-size:12px;color:var(--suave);cursor:pointer}.st b{display:block;font-size:20px;color:var(--tinta)}.st.on{outline:2px solid var(--tinta)}
input.q{border:1px solid var(--linha);border-radius:10px;padding:8px 12px;font-size:13px;min-width:260px;background:#fff}
.lei{font-size:11.5px;color:var(--suave);max-width:520px}.lei details summary{cursor:pointer;color:var(--verde);font-weight:600}
.foot{font-size:11px;color:var(--fraco);margin-top:14px;line-height:1.5}
.sim td input{width:70px;border:1px solid var(--linha);border-radius:6px;padding:4px 6px;text-align:right}
.vazio{padding:30px;text-align:center;color:var(--fraco);background:#fff;border-radius:14px;border:1px dashed var(--linha)}
.mais{margin-top:8px;font-size:12px;color:var(--fraco)}
@media (max-width:900px){.bento{grid-template-columns:1fr}}
</style></head><body><div class="wrap">
<div class="hd"><div><h1>Apuração Assistida<span class="rt">REFORMA TRIBUTÁRIA</span></h1></div><div class="m" id="cab"></div></div>
<div class="tabs" id="tabs"><span data-p="painel" class="on">Painel</span><span data-p="apur">Apuração</span><span data-p="clas">Classificação CST / cClassTrib</span><span data-p="pont">Pontos de atenção</span><span data-p="sim">Simulação até 2033</span><span data-p="xs">XML × SPED</span></div>
<div class="pg on" id="pg-painel"></div><div class="pg" id="pg-apur"></div><div class="pg" id="pg-clas"></div><div class="pg" id="pg-pont"></div><div class="pg" id="pg-sim"></div><div class="pg" id="pg-xs"></div>
<div class="foot" id="foot"></div></div>
<script>
const D=/*DADOS*/;
const $=s=>document.querySelector(s);
const m=v=>(v||0).toLocaleString('pt-BR',{minimumFractionDigits:2,maximumFractionDigits:2});
const e=s=>String(s==null?'':s).replace(/[&<>"]/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;'}[c]));
const br=d=>d&&d.length>=10?d.slice(8,10)+'/'+d.slice(5,7)+'/'+d.slice(0,4):(d||'');
const tg=s=>`<span class="tg t-${s.split(' ')[0].replace('/','')}">${e(s)}</span>`;
const tp=p=>p==='debito'?'Saída':'Entrada';
const P=D.painel;
$('#cab').innerHTML=`${e(D.empresa||'')} · CNPJ ${e(D.cnpj)}<br>competência ${D.meses.map(x=>x.slice(5)+'/'+x.slice(0,4)).join(', ')} · ${D.notas.debito} notas de saída · ${D.notas.credito} de entrada`+(D.regime?`<br>regime específico informado: ${e(D.regime_nome)} (${e(D.regime)})`:'');
$('#foot').innerHTML=`Base legal: LC 214/2025 — art. 46 (apuração assistida), art. 47 (créditos), art. 57 (uso ou consumo pessoal), art. 343 e seguintes (ano de teste 2026). Classificação conferida com a tabela oficial cClassTrib ${e(D.tabela.versao)} (atualizada em ${br(D.tabela.atualizacao)}; ${D.tabela.codigos} códigos, ${D.tabela.revogados} com vigência encerrada). ${D.xmls_lidos} arquivos lidos · ${D.canceladas} cancelamento(s) no lote${Object.entries(D.ignorados).map(([k,v])=>` · ${v} ${k}`).join('')}.`;
document.querySelectorAll('#tabs span').forEach(t=>t.onclick=()=>{document.querySelectorAll('#tabs span').forEach(x=>x.classList.toggle('on',x===t));document.querySelectorAll('.pg').forEach(x=>x.classList.toggle('on',x.id==='pg-'+t.dataset.p));});
// ---------- Painel
(function(){
 const deb=P.debito.total,cred=P.credito.total,sal=P.saldo.total,H=250,y0=30;
 const hc=deb>0?Math.max(6,Math.min(H-6,H*Math.min(cred,deb)/deb)):H/2,hs=H-hc,pct=deb>0?Math.round(100*Math.min(cred,deb)/deb):0;
 const fl=`<svg width="100%" viewBox="0 0 470 320"><rect x="10" y="${y0}" width="120" height="${H}" rx="10" fill="#e2552d"/>
 <text x="70" y="${y0+H/2-6}" fill="#fff" font-size="13" text-anchor="middle" font-weight="700">Débitos</text><text x="70" y="${y0+H/2+16}" fill="#fff" font-size="16" text-anchor="middle" font-weight="800">${m(deb)}</text>
 <path d="M130 ${y0} C200 ${y0},200 ${y0+8},270 ${y0+8} L270 ${y0+8+hc} C200 ${y0+8+hc},200 ${y0+hc},130 ${y0+hc} Z" fill="#0f766e" opacity=".85"/>
 <path d="M130 ${y0+hc} C200 ${y0+hc},200 ${y0+hc+24},270 ${y0+hc+24} L270 ${y0+H+24} C200 ${y0+H+24},200 ${y0+H},130 ${y0+H} Z" fill="#1c1917" opacity=".9"/>
 <rect x="270" y="${y0+8}" width="190" height="${hc}" rx="10" fill="#0f766e"/>
 <text x="365" y="${y0+8+hc/2-6}" fill="#fff" font-size="12.5" text-anchor="middle" font-weight="700">compensado por créditos</text><text x="365" y="${y0+8+hc/2+16}" fill="#fff" font-size="17" text-anchor="middle" font-weight="800">${m(cred)}</text>
 <rect x="270" y="${y0+hc+24}" width="190" height="${Math.max(hs,40)}" rx="10" fill="#1c1917"/>
 <text x="365" y="${y0+hc+24+Math.max(hs,40)/2-4}" fill="#fff" font-size="12.5" text-anchor="middle" font-weight="700">${sal>=0?'saldo a recolher':'saldo credor'}</text><text x="365" y="${y0+hc+24+Math.max(hs,40)/2+16}" fill="#fbbf24" font-size="17" text-anchor="middle" font-weight="800">${m(Math.abs(sal))}</text></svg>`;
 const tot=Math.abs(P.saldo.cbs)+Math.abs(P.saldo.ibs)||1,fc=100*Math.abs(P.saldo.cbs)/tot;
 const S=D.situacao,nd=S.Divergente||0,nv=S.Verificar||0;
 const cont={};D.pontos.forEach(p=>{const k=p.texto.split(/[:—(]/)[0].trim().slice(0,52);cont[k]=(cont[k]||0)+1});
 const chips=Object.entries(cont).sort((a,b)=>b[1]-a[1]).slice(0,5).map(([k,v])=>`<span class="chip">${v} · ${e(k)}</span>`).join('');
 // barras da transição
 const b=D.simulacao.bases,cen=D.simulacao.cenarios,vals=cen.map(c=>({a:c.ano,cbs:((b.debito_cbs||0)-(b.credito_cbs||0))*c.cbs/100,ibs:((b.debito_ibs||0)-(b.credito_ibs||0))*c.ibs/100}));
 const mx=Math.max(...vals.map(v=>Math.max(0,v.cbs)+Math.max(0,v.ibs)),1);
 const bars=vals.map((v,i)=>{const x=20+i*75,hc=110*Math.max(0,v.cbs)/mx,hi=110*Math.max(0,v.ibs)/mx,top=150-hc-hi;const k=v.cbs+v.ibs;
  return `<rect x="${x}" y="${150-hc}" width="48" height="${hc}" fill="#e2552d"/><rect x="${x}" y="${top}" width="48" height="${hi}" fill="#0f766e"/><text x="${x+24}" y="167" text-anchor="middle" font-size="10.5" fill="#57534e">${v.a}</text><text x="${x+24}" y="${top-6}" text-anchor="middle" font-size="10.5" font-weight="700" fill="#1c1917">${k>=1000?(k/1000).toLocaleString('pt-BR',{maximumFractionDigits:1})+' mil':m(k)}</text>`}).join('');
 $('#pg-painel').innerHTML=`<div class="bento">
 <div class="c" style="grid-row:span 2"><h3>Do débito ao saldo</h3><div class="lbl">IBS + CBS destacados nos XMLs do período</div>${fl}
 <div class="legend"><span><i style="background:#e2552d"></i>saídas</span><span><i style="background:#0f766e"></i>créditos das entradas (${pct}% do débito)</span><span><i style="background:#1c1917"></i>saldo</span></div>
 <div style="margin-top:10px"><div class="aj"><span>Crédito estimado (fornecedor do regime normal sem IBS/CBS no XML)</span><b>${m(P.estimado.total)}</b></div>
 <div class="aj"><span>Crédito a confirmar — art. 57 (uso pessoal × atividade)</span><b>${m(P.art57.total)}</b></div>
 <div class="aj"><span>Saldo se os dois forem confirmados</span><b>${m(P.saldo_com_ajustes)}</b></div>
 ${P.simples.notas?`<div class="aj"><span>${P.simples.notas} nota(s) de fornecedor do Simples sem IBS/CBS — sem crédito</span><b></b></div>`:''}</div></div>
 <div class="c"><h3>Saldo por tributo</h3><div style="display:flex;align-items:center;gap:14px"><svg width="118" height="118" viewBox="0 0 42 42"><circle cx="21" cy="21" r="15.9" fill="none" stroke="#f4f1ea" stroke-width="6"/><circle cx="21" cy="21" r="15.9" fill="none" stroke="#e2552d" stroke-width="6" stroke-dasharray="${fc} ${100-fc}" stroke-dashoffset="25"/><circle cx="21" cy="21" r="15.9" fill="none" stroke="#0f766e" stroke-width="6" stroke-dasharray="${100-fc} ${fc}" stroke-dashoffset="${25-fc}"/></svg>
 <div><div class="lbl">CBS</div><div style="font-size:21px;font-weight:800;color:#e2552d">${m(P.saldo.cbs)}</div><div class="lbl" style="margin-top:6px">IBS (UF ${m(P.debito.ibsuf-P.credito.ibsuf)} · município ${m(P.debito.ibsmun-P.credito.ibsmun)})</div><div style="font-size:21px;font-weight:800;color:#0f766e">${m(P.saldo.ibs)}</div></div></div></div>
 <div class="c dark"><h3>Pontos de atenção</h3><div class="big" style="color:#fbbf24">${nd+nv}</div><div style="font-size:12.5px;color:#d6d3d1">${nd} itens divergentes · ${nv} para verificar · ${S.Confere||0} conferem</div><div style="margin-top:6px">${chips}</div></div>
 <div class="c" style="grid-column:span 2"><h3>Transição até 2033 — simulação com a mesma movimentação</h3><svg width="100%" viewBox="0 0 620 175">${bars}</svg>
 <div class="legend"><span><i style="background:#e2552d"></i>CBS</span><span><i style="background:#0f766e"></i>IBS</span><span>alíquotas na aba Simulação · não são valores oficiais</span></div></div></div>`;
})();
// ---------- Apuração
(function(){
 const lin=D.por_cfop.map(x=>`<tr><td>${tp(x.papel)}</td><td class="code">${e(x.cfop)}</td><td class="n">${x.itens}</td><td class="n">${m(x.bc)}</td><td class="n">${m(x.ibs)}</td><td class="n">${m(x.cbs)}</td></tr>`).join('');
 const docs=D.documentos.slice(0,1500).map(x=>`<tr><td>${tp(x.papel)}</td><td>${x.mod==='65'?'NFC-e':'NF-e'} ${e(x.num)}</td><td>${br(x.data)}</td><td>${e(x.part)}<div class="sm">${e(x.part_doc)}</div></td><td class="n">${m(x.vnf)}</td><td class="n">${m(x.bc)}</td><td class="n">${m(x.ibs)}</td><td class="n">${m(x.cbs)}</td><td>${tg(x.sit)}</td></tr>`).join('');
 $('#pg-apur').innerHTML=`<h3 style="margin:0 0 8px">Por CFOP</h3><table><tr><th>Tipo</th><th>CFOP</th><th class="n">Itens</th><th class="n">Base IBS/CBS</th><th class="n">IBS</th><th class="n">CBS</th></tr>${lin}</table>
 <h3 style="margin:18px 0 8px">Por documento</h3><table><tr><th>Tipo</th><th>Nota</th><th>Data</th><th>Participante</th><th class="n">Valor</th><th class="n">Base</th><th class="n">IBS</th><th class="n">CBS</th><th>Situação</th></tr>${docs}</table>
 ${D.documentos.length>1500?`<div class="mais">Mostrando 1.500 de ${D.documentos.length} documentos — a lista completa está no Excel.</div>`:''}`;
})();
// ---------- Classificação
(function(){
 let filtro='',busca='';
 const S={};D.produtos.forEach(p=>S[p.sit]=(S[p.sit]||0)+1);
 const cards=[['','Todos',D.produtos.length],['Confere','Conferem com a tabela oficial',S.Confere||0],['Divergente','Divergentes',S.Divergente||0],['Verificar','Para verificar',S.Verificar||0],['Sem IBS/CBS','Sem IBS/CBS no XML',S['Sem IBS/CBS']||0]];
 const cls=D.classes.map(c=>`<tr><td class="code">${e(c.cst)}</td><td class="code">${e(c.cclass)}</td><td>${e(c.nome)}</td><td class="n">${c.red?c.red+'%':'—'}</td><td>${e(c.artigo)}</td><td class="n">${c.itens}</td></tr>`).join('');
 const desenha=()=>{const l=D.produtos.filter(p=>(!filtro||p.sit===filtro)&&(!busca||(p.desc+' '+p.ncm+' '+p.cod+' '+p.cclass).toLowerCase().includes(busca))).slice(0,800);
  $('#clasT').innerHTML=l.map(p=>`<tr><td>${e(p.desc)}<div class="sm">${tp(p.papel)} · cód. ${e(p.cod)} · ${p.itens} item(ns) em ${p.notas} nota(s)</div></td><td class="code">${e(p.ncm)}</td>
  <td>${p.cclass?`<span class="code">CST ${e(p.cst)} · ${e(p.cclass)}</span><div class="sm">${e(p.nome)}</div>`:'<span class="sm">sem grupo IBS/CBS</span>'}</td>
  <td class="lei">${p.motivos.map(x=>`<div>• ${e(x.texto)}${x.qtd>1?` <span class="sm">(${x.qtd})</span>`:''}</div>`).join('')}${p.texto?`<details><summary>${e(p.artigo)} — texto da lei</summary>${e(p.texto)}</details>`:''}</td>
  <td class="n">${m(p.ibs+p.cbs)}</td><td>${tg(p.sit)}</td></tr>`).join('')||'<tr><td colspan="6" class="sm">Nada com esse filtro.</td></tr>';};
 $('#pg-clas').innerHTML=`<div class="bar">${cards.map(([k,t,n])=>`<div class="st${k===''?' on':''}" data-f="${k}"><b>${n}</b>${t}</div>`).join('')}<input class="q" placeholder="Buscar produto, NCM ou cClassTrib"></div>
 <table><tr><th>Produto</th><th>NCM</th><th>No XML</th><th>Conferência com a tabela oficial</th><th class="n">IBS+CBS</th><th>Situação</th></tr><tbody id="clasT"></tbody></table>
 <h3 style="margin:18px 0 8px">Classificações usadas no período</h3><table><tr><th>CST</th><th>cClassTrib</th><th>Classificação (tabela oficial)</th><th class="n">Redução</th><th>Artigo</th><th class="n">Itens</th></tr>${cls}</table>`;
 document.querySelectorAll('#pg-clas .st').forEach(c=>c.onclick=()=>{filtro=c.dataset.f;document.querySelectorAll('#pg-clas .st').forEach(x=>x.classList.toggle('on',x===c));desenha();});
 $('#pg-clas .q').oninput=ev=>{busca=ev.target.value.toLowerCase();desenha();};desenha();
})();
// ---------- Pontos
(function(){
 if(!D.pontos.length){$('#pg-pont').innerHTML='<div class="vazio">Nenhum ponto de atenção.</div>';return;}
 $('#pg-pont').innerHTML=`<table><tr><th>Situação</th><th>Nota</th><th>Produto</th><th>CST / cClassTrib</th><th>O que foi encontrado</th></tr>${D.pontos.slice(0,2000).map(p=>`<tr><td>${tg(p.sit)}</td><td>${tp(p.papel)} ${p.mod==='65'?'NFC-e':'NF-e'} ${e(p.num)}<div class="sm">${br(p.data)} · item ${e(p.item)}</div></td><td>${e(p.desc)}<div class="sm">NCM ${e(p.ncm)}</div></td><td class="code">${e(p.cst)} · ${e(p.cclass)}</td><td>${e(p.texto)}</td></tr>`).join('')}</table>
 ${D.n_pontos>2000?`<div class="mais">Mostrando 2.000 de ${D.n_pontos} — todos estão no Excel.</div>`:''}
 ${D.simples_sem_grupo.length?`<h3 style="margin:18px 0 8px">Fornecedores do Simples Nacional sem IBS/CBS no XML (sem crédito)</h3><table><tr><th>Nota</th><th>Data</th><th>Fornecedor</th><th>Regime</th><th class="n">Valor</th></tr>${D.simples_sem_grupo.map(x=>`<tr><td>${e(x.num)}</td><td>${br(x.data)}</td><td>${e(x.nome)}<div class="sm">${e(x.emit)}</div></td><td>${e(x.regime)}</td><td class="n">${m(x.vnf)}</td></tr>`).join('')}</table>`:''}`;
})();
// ---------- Simulação
(function(){
 const b=D.simulacao.bases,cen=JSON.parse(JSON.stringify(D.simulacao.cenarios));
 const desenha=()=>{$('#simT').innerHTML=cen.map((c,i)=>{const deb=(b.debito_cbs||0)*c.cbs/100+(b.debito_ibs||0)*c.ibs/100,cred=(b.credito_cbs||0)*c.cbs/100+(b.credito_ibs||0)*c.ibs/100;
  return `<tr><td><b>${c.ano}</b></td><td class="n"><input data-i="${i}" data-k="cbs" value="${String(c.cbs).replace('.',',')}">%</td><td class="n"><input data-i="${i}" data-k="ibs" value="${String(c.ibs).replace('.',',')}">%</td><td class="n">${m(deb)}</td><td class="n">${m(cred)}</td><td class="n"><b>${m(deb-cred)}</b></td><td class="sm">${e(c.nota)}</td></tr>`}).join('');
  document.querySelectorAll('#simT input').forEach(inp=>inp.onchange=()=>{const v=parseFloat(inp.value.replace(',','.'));if(!isNaN(v)){cen[inp.dataset.i][inp.dataset.k]=v;desenha();}});};
 $('#pg-sim').innerHTML=`<div class="c" style="margin-bottom:12px;font-size:12.5px;color:#57534e">Mesma movimentação do período, com as reduções de cada cClassTrib, nas alíquotas de cada ano. Entram os valores destacados nos XMLs (sem os créditos estimados e os do art. 57). As alíquotas podem ser alteradas aqui. Em 2026 pode haver diferença de centavos para o Painel porque os XMLs arredondam item por item. <b>Não são valores oficiais</b>: a CBS de 2027 (11,3%) foi calculada pelo TCU e depende de resolução do Senado, e o IBS de referência ainda não foi fixado. O ICMS em redução de 2029 a 2032 não entra na conta.</div>
 <table class="sim"><tr><th>Ano</th><th class="n">CBS</th><th class="n">IBS</th><th class="n">Débitos</th><th class="n">Créditos</th><th class="n">Saldo</th><th>Observação</th></tr><tbody id="simT"></tbody></table>`;desenha();
})();
// ---------- XML × SPED
(function(){
 const x=D.xml_sped;
 if(!x){$('#pg-xs').innerHTML='<div class="vazio">Nenhum SPED Fiscal foi anexado. Anexe o SPED (opcional) para ver: XML sem escrituração, nota escriturada sem XML no lote e nota cancelada no SPED que ainda está autorizada no lote.</div>';return;}
 const t=(tit,l)=>`<h3 style="margin:16px 0 8px">${tit} <span class="sm">(${l.length})</span></h3>`+(l.length?`<table><tr><th>Tipo</th><th>Nota</th><th>Data</th><th class="n">Valor</th><th>Chave</th></tr>${l.slice(0,1000).map(k=>`<tr><td>${e(k.tipo)}</td><td>${e(k.num)}</td><td>${br(k.data)}</td><td class="n">${m(k.vnf)}</td><td class="code" style="font-size:11px">${e(k.chave)}</td></tr>`).join('')}</table>`:'<div class="sm">Nenhuma.</div>');
 $('#pg-xs').innerHTML=`<div class="sm">SPED ${e(x.empresa)} · ${e(x.periodo)}${x.cnpj_difere?' · <b style="color:#b91c1c">o CNPJ do SPED é diferente do CNPJ informado</b>':''}</div>`+t('XML sem escrituração no SPED',x.sem_sped)+t('Escriturada no SPED sem XML no lote',x.sem_xml)+t('Cancelada no SPED e autorizada no lote',x.canc_pasta);
})();
</script></body></html>'''


# ------------------------------------------------------------------------------------------------------------ CLI
def main_cli(argv):
    def opt(n):
        return argv[argv.index(n) + 1] if n in argv and argv.index(n) + 1 < len(argv) else None

    def lista(n):
        if n not in argv:
            return []
        out = []
        for a in argv[argv.index(n) + 1:]:
            if a.startswith('--'):
                break
            out.append(a)
        return out
    xmls, js = lista('--xmls'), opt('--json')
    if not xmls or not js:
        print('uso: apuracao-rtc --xmls a.zip [b.zip ...] --json saida.json [--sped arq.txt] [--cnpj N] [--regime cClassTrib] '
              '[--html saida.html] [--xlsx saida.xlsx]')
        return 1
    try:
        res = analisar(xmls, cnpj=opt('--cnpj') or '', regime=opt('--regime') or '', sped=opt('--sped') or '')
        if not res['notas']['debito'] and not res['notas']['credito']:
            res = {'erro': 'Nenhuma NF-e ou NFC-e do CNPJ foi encontrada nos arquivos selecionados.'}
        else:
            if opt('--html'):
                tmp = opt('--html') + '.parcial'
                with open(tmp, 'w', encoding='utf-8') as fh:
                    fh.write(gerar_html(res))
                os.replace(tmp, opt('--html'))
            if opt('--xlsx'):
                gerar_xlsx(res, opt('--xlsx'))
            res = {k: v for k, v in res.items() if k not in ('itens', 'documentos', 'pontos', 'produtos')} | {
                'n_pontos': len(res['pontos']), 'n_produtos': len(res['produtos'])}
    except Exception as e:  # pragma: no cover
        import traceback
        traceback.print_exc()
        res = {'erro': f'Falha ao processar: {e}'}
    tmp = js + '.parcial'
    with open(tmp, 'w', encoding='utf-8') as fh:
        json.dump(res, fh, ensure_ascii=False)
    os.replace(tmp, js)
    return 0
