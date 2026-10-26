# -*- coding: utf-8 -*-
"""Consulta Manifestação NF-e — situação da nota na SEFAZ e manifestação do destinatário, pela chave de acesso.

SÓ CONSULTA: nenhum evento é enviado. Para cada chave:
  1. NfeConsultaProtocolo 4.00 na SEFAZ que autorizou a nota (a UF está nos 2 primeiros dígitos da chave):
     situação (autorizada / cancelada / denegada) e os eventos registrados — entre eles a manifestação
     do destinatário (210200 Confirmação, 210210 Ciência, 210220 Desconhecimento, 210240 Operação não realizada).
  2. NFeDistribuicaoDFe (Ambiente Nacional), consulta por chave: o resumo da nota (emitente, valor, data de emissão),
     quando o Ambiente Nacional devolve. Se não devolver, esses campos ficam em branco — nada é estimado.
O certificado A1 usado é o da empresa destinatária (o mesmo cadastrado nas Baixas DF-e).
Prazo do evento conclusivo: 180 dias contados da autorização da nota (NT 2020.001 da NF-e).
"""
import base64
import gzip
import json
import re
import time
from datetime import datetime, timezone

UF_SIGLA = {'11': 'RO', '12': 'AC', '13': 'AM', '14': 'RR', '15': 'PA', '16': 'AP', '17': 'TO', '21': 'MA', '22': 'PI',
            '23': 'CE', '24': 'RN', '25': 'PB', '26': 'PE', '27': 'AL', '28': 'SE', '29': 'BA', '31': 'MG', '32': 'ES',
            '33': 'RJ', '35': 'SP', '41': 'PR', '42': 'SC', '43': 'RS', '50': 'MS', '51': 'MT', '52': 'GO', '53': 'DF'}

# NfeConsultaProtocolo 4.00 — produção. SP, MG e SVRS conferidos nos portais oficiais (26/10/2026); os demais conforme a
# relação de serviços do Portal Nacional da NF-e. Se o primeiro endereço falhar, tenta os seguintes.
_WS = {
    'SVRS': 'https://nfe.svrs.rs.gov.br/ws/NfeConsulta/NfeConsulta4.asmx',
    'SVAN': 'https://www.sefazvirtual.fazenda.gov.br/NFeConsultaProtocolo4/NFeConsultaProtocolo4.asmx',
    'SP': 'https://nfe.fazenda.sp.gov.br/ws/nfeconsultaprotocolo4.asmx',
    'MG': 'https://nfe.fazenda.mg.gov.br/nfe2/services/NFeConsultaProtocolo4',
    'AM': 'https://nfe.sefaz.am.gov.br/services2/services/NfeConsulta4',
    'BA': 'https://nfe.sefaz.ba.gov.br/webservices/NFeConsultaProtocolo4/NFeConsultaProtocolo4.asmx',
    'CE': 'https://nfe.sefaz.ce.gov.br/nfe4/services/NFeConsultaProtocolo4',
    'GO': 'https://nfe.sefaz.go.gov.br/nfe/services/NFeConsultaProtocolo4',
    'MS': 'https://nfe.sefaz.ms.gov.br/ws/NFeConsultaProtocolo4',
    'MT': 'https://nfe.sefaz.mt.gov.br/nfews/v2/services/NfeConsulta4',
    'PE': 'https://nfe.sefaz.pe.gov.br/nfe-service/services/NFeConsultaProtocolo4',
    'PR': 'https://nfe.sefa.pr.gov.br/nfe/NFeConsultaProtocolo4',
    'RS': 'https://nfe.sefazrs.rs.gov.br/ws/NfeConsulta/NfeConsulta4.asmx',
}
_PROPRIO = {'13': 'AM', '23': 'CE', '26': 'PE', '29': 'BA', '31': 'MG', '35': 'SP', '41': 'PR', '43': 'RS', '50': 'MS', '51': 'MT', '52': 'GO'}
_URL_DIST = 'https://www1.nfe.fazenda.gov.br/NFeDistribuicaoDFe/NFeDistribuicaoDFe.asmx'

EVENTOS = {'110111': 'cancelamento', '210200': 'confirmada', '210210': 'ciencia', '210220': 'desconhecida', '210240': 'nao_realizada'}
CONCLUSIVOS = ('confirmada', 'desconhecida', 'nao_realizada')
PRAZO_DIAS = 180


def chave_valida(ch):
    ch = re.sub(r'\D', '', str(ch or ''))
    if len(ch) != 44:
        return False
    soma, peso = 0, 2
    for d in reversed(ch[:43]):
        soma += int(d) * peso
        peso = 2 if peso == 9 else peso + 1
    dv = 11 - soma % 11
    return int(ch[43]) == (0 if dv >= 10 else dv)


def extrair_chaves(texto):
    """Chaves de 44 dígitos num texto qualquer (aceita espaços e pontos no meio de cada linha)."""
    vistos, ok, ruins = set(), [], []
    for i, linha in enumerate(str(texto or '').splitlines(), 1):
        achou = re.findall(r'(?<!\d)\d{44}(?!\d)', re.sub(r'[\s.\-/]', '', linha))
        if not achou:
            dig = re.sub(r'\D', '', linha)
            if dig:
                ruins.append({'linha': i, 'texto': linha.strip()[:60], 'motivo': f'{len(dig)} dígitos (a chave tem 44)'})
            continue
        for ch in achou:
            if ch in vistos:
                continue
            vistos.add(ch)
            (ok if chave_valida(ch) else ruins).append(ch if chave_valida(ch) else {'linha': i, 'texto': ch, 'motivo': 'dígito verificador não confere'})
    return ok, ruins


def dados_da_chave(ch):
    return {'uf': UF_SIGLA.get(ch[:2], ch[:2]), 'aamm': f'{ch[4:6]}/20{ch[2:4]}', 'cnpj_emit': ch[6:20], 'modelo': ch[20:22],
            'serie': str(int(ch[22:25])), 'numero': str(int(ch[25:34]))}


def _txt(xml, tag):
    m = re.search(rf'<(?:\w+:)?{tag}\b[^>]*>([^<]*)</(?:\w+:)?{tag}>', xml)
    return m.group(1).strip() if m else ''


def _blocos(xml, tag):
    return re.findall(rf'<(?:\w+:)?{tag}\b[^>]*>.*?</(?:\w+:)?{tag}>', xml, re.S)


def _dt(s):
    try:
        return datetime.fromisoformat(s.replace('Z', '+00:00'))
    except (ValueError, AttributeError):
        return None


def _soap(sess, url, ns_wsdl, corpo, acao, timeout=30):
    env = ('<?xml version="1.0" encoding="utf-8"?><soap12:Envelope xmlns:soap12="http://www.w3.org/2003/05/soap-envelope">'
           f'<soap12:Body><nfeDadosMsg xmlns="{ns_wsdl}">{corpo}</nfeDadosMsg></soap12:Body></soap12:Envelope>')
    r = sess.post(url, data=env.encode('utf-8'), timeout=timeout,
                  headers={'Content-Type': f'application/soap+xml; charset=utf-8; action="{ns_wsdl}/{acao}"'})
    r.raise_for_status()
    return r.text


def _consulta_situacao(sess, ch):
    corpo = (f'<consSitNFe xmlns="http://www.portalfiscal.inf.br/nfe" versao="4.00"><tpAmb>1</tpAmb>'
             f'<xServ>CONSULTAR</xServ><chNFe>{ch}</chNFe></consSitNFe>')
    ns = 'http://www.portalfiscal.inf.br/nfe/wsdl/NFeConsultaProtocolo4'
    prim = _PROPRIO.get(ch[:2], 'SVAN' if ch[:2] == '21' else 'SVRS')
    ordem = [prim] + [a for a in ('SVRS', 'SVAN') if a != prim]
    ultimo = None
    for aut in ordem:
        try:
            xml = _soap(sess, _WS[aut], ns, corpo, 'nfeConsultaNF')
        except Exception as e:          # endereço fora do ar ou de outro autorizador: tenta o próximo
            ultimo = f'{aut}: {str(e)[:140]}'
            continue
        cst = _txt(xml, 'cStat')
        if cst in ('226', '227') and aut != ordem[-1]:   # UF da chave não é desse autorizador
            ultimo = f'{aut}: {_txt(xml, "xMotivo")}'
            continue
        return xml, aut
    raise RuntimeError(ultimo or 'Sem resposta da SEFAZ.')


def _analisa_situacao(xml):
    ret = re.search(r'<(?:\w+:)?retConsSitNFe\b.*?</(?:\w+:)?retConsSitNFe>', xml, re.S)
    ret = ret.group(0) if ret else xml
    cst, mot = _txt(ret, 'cStat'), _txt(ret, 'xMotivo')
    prot = _blocos(ret, 'protNFe')
    dh_aut = _txt(prot[0], 'dhRecbto') if prot else ''
    cst_prot = _txt(prot[0], 'cStat') if prot else ''
    eventos = []
    for pe in _blocos(ret, 'procEventoNFe'):
        tp = _txt(pe, 'tpEvento')
        rets = _blocos(pe, 'retEvento')
        st_ev = _txt(rets[0], 'cStat') if rets else ''
        if st_ev and st_ev not in ('135', '136', '155'):
            continue
        dh = (_txt(rets[0], 'dhRegEvento') if rets else '') or _txt(pe, 'dhEvento')
        eventos.append({'tipo': tp, 'nome': EVENTOS.get(tp, tp), 'quando': dh})
    if cst == '101' or cst_prot == '101' or any(e['tipo'] == '110111' for e in eventos):
        sit = 'cancelada'
    elif cst in ('110', '301', '302', '303') or cst_prot in ('110', '301', '302', '303'):
        sit = 'denegada'
    elif cst == '100' or cst_prot == '100':
        sit = 'autorizada'
    elif cst == '217':
        sit = 'nao_encontrada'
    else:
        sit = 'erro'
    manif = [e for e in eventos if e['nome'] in ('ciencia',) + CONCLUSIVOS]
    concl = sorted((e for e in manif if e['nome'] in CONCLUSIVOS), key=lambda e: e['quando'])
    cien = sorted((e for e in manif if e['nome'] == 'ciencia'), key=lambda e: e['quando'])
    if concl:
        m, dh_m = concl[-1]['nome'], concl[-1]['quando']
    elif cien:
        m, dh_m = 'ciencia', cien[0]['quando']
    else:
        m, dh_m = 'nenhuma', ''
    r = {'situacao': sit, 'cstat': cst, 'motivo': mot, 'dh_autorizacao': dh_aut, 'manifestacao': m, 'dh_manifestacao': dh_m,
         'eventos': eventos}
    d = _dt(dh_aut)
    if d and m in ('nenhuma', 'ciencia') and sit == 'autorizada':
        passados = (datetime.now(timezone.utc) - d.astimezone(timezone.utc)).days
        r['prazo_dias'] = PRAZO_DIAS - passados
    return r


def _resumo_an(sess, ch, cnpj, cuf):
    corpo = (f'<nfeDistDFeInteresse xmlns="http://www.portalfiscal.inf.br/nfe/wsdl/NFeDistribuicaoDFe"><nfeDadosMsg>'
             f'<distDFeInt xmlns="http://www.portalfiscal.inf.br/nfe" versao="1.01"><tpAmb>1</tpAmb><cUFAutor>{cuf}</cUFAutor>'
             f'<CNPJ>{cnpj}</CNPJ><consChNFe><chNFe>{ch}</chNFe></consChNFe></distDFeInt></nfeDadosMsg></nfeDistDFeInteresse>')
    env = ('<?xml version="1.0" encoding="utf-8"?><soap12:Envelope xmlns:soap12="http://www.w3.org/2003/05/soap-envelope">'
           f'<soap12:Body>{corpo}</soap12:Body></soap12:Envelope>')
    r = sess.post(_URL_DIST, data=env.encode('utf-8'), timeout=30, headers={
        'Content-Type': 'application/soap+xml; charset=utf-8; action="http://www.portalfiscal.inf.br/nfe/wsdl/NFeDistribuicaoDFe/nfeDistDFeInteresse"'})
    r.raise_for_status()
    xml = r.text
    cst = _txt(xml, 'cStat')
    out = {'cstat_an': cst}
    for dz in re.findall(r'<(?:\w+:)?docZip\b[^>]*>([^<]+)</(?:\w+:)?docZip>', xml):
        try:
            doc = gzip.decompress(base64.b64decode(dz)).decode('utf-8', 'replace')
        except Exception:
            continue
        if '<resNFe' in doc or '<nfeProc' in doc or '<NFe' in doc:
            emit = _blocos(doc, 'emit')
            out['emitente'] = _txt(emit[0], 'xNome') if emit else _txt(doc, 'xNome')
            out['valor'] = _txt(doc, 'vNF')
            out['dh_emissao'] = _txt(doc, 'dhEmi')
            break
    return out


def consultar(chaves, cert, key, cnpj, cuf='11', progresso=None, com_resumo=True, pausa=0.35):
    """cert/key: PEM temporários do certificado A1 (quem chama extrai do .pfx e apaga depois)."""
    import requests
    sess = requests.Session()
    sess.cert = (cert, key)
    res, an_ok = [], com_resumo
    try:
        for i, ch in enumerate(chaves, 1):
            item = {'chave': ch, **dados_da_chave(ch)}
            try:
                xml, aut = _consulta_situacao(sess, ch)
                item.update(_analisa_situacao(xml)); item['autorizador'] = aut
            except Exception as e:
                item.update({'situacao': 'erro', 'motivo': str(e)[:200], 'manifestacao': 'nenhuma'})
            if an_ok and item.get('situacao') in ('autorizada', 'cancelada', 'denegada'):
                try:
                    r = _resumo_an(sess, ch, cnpj, cuf)
                    if r.get('cstat_an') == '656':        # consumo indevido: não insiste no resto do lote
                        an_ok = False
                    item.update({k: v for k, v in r.items() if k != 'cstat_an' and v})
                except Exception:
                    pass
            res.append(item)
            if progresso:
                try:
                    with open(progresso, 'w', encoding='utf-8') as fh:
                        json.dump({'feitas': i, 'total': len(chaves)}, fh)
                except OSError:
                    pass
            time.sleep(pausa)
    finally:
        sess.close()
    return res


def exportar_xlsx(itens, destino, titulo=''):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    nome_m = {'confirmada': 'Confirmada', 'ciencia': 'Só ciência', 'nenhuma': 'Sem manifestação', 'desconhecida': 'Desconhecida',
              'nao_realizada': 'Operação não realizada'}
    nome_s = {'autorizada': 'Autorizada', 'cancelada': 'Cancelada', 'denegada': 'Denegada', 'nao_encontrada': 'Não encontrada', 'erro': 'Erro na consulta'}
    wb = Workbook(); ws = wb.active; ws.title = 'Manifestação'
    if titulo:
        ws.append([titulo]); ws['A1'].font = Font(bold=True, size=12); ws.append([])
    cab = ['Chave de acesso', 'UF', 'CNPJ emitente', 'Emitente', 'Nº', 'Série', 'Emissão', 'Valor (R$)', 'Situação', 'Manifestação',
           'Data da manifestação', 'Prazo (dias)', 'Observação']
    ws.append(cab)
    lin_cab = ws.max_row
    for c in ws[lin_cab]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1E3A8A'); c.alignment = Alignment(vertical='center')
    for it in itens:
        try:
            v = float(it.get('valor') or '')
        except ValueError:
            v = None
        ws.append([it['chave'], it.get('uf', ''), it.get('cnpj_emit', ''), it.get('emitente', ''), it.get('numero', ''), it.get('serie', ''),
                   (it.get('dh_emissao') or '')[:10] or it.get('aamm', ''), v, nome_s.get(it.get('situacao'), it.get('situacao', '')),
                   nome_m.get(it.get('manifestacao'), ''), (it.get('dh_manifestacao') or '')[:10], it.get('prazo_dias', ''),
                   it.get('motivo', '') if it.get('situacao') in ('erro', 'nao_encontrada') else ''])
        ws.cell(ws.max_row, 8).number_format = '#,##0.00'
    for col, w in zip('ABCDEFGHIJKLM', (48, 5, 17, 36, 10, 6, 12, 14, 16, 22, 18, 12, 40)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = ws.cell(lin_cab + 1, 1)
    wb.save(destino)
