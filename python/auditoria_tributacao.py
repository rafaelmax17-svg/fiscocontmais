# -*- coding: utf-8 -*-
"""Análise Tributária (menu Fiscal, equipe) — Tributação das entradas por produto + auditoria de tributação.

Módulo de TESTE do Admin (pedido do Rafael, 29/10/2026). Depois de validado, vai para o Auditor de Classificação.
Só lê o SPED Fiscal (EFD ICMS/IPI): nada é alterado. Tudo é apontamento para o contador conferir.

Fonte dos dados: C100 de ENTRADA (IND_OPER = 0, sem cancelados) + C170 (itens) + 0200 (produto/NCM/tipo) + 0150 (fornecedor)
+ 0190 (unidade) + C197 (ajustes da nota). Nota de entrada sem C170 entra pelos totais do C190, marcada "nota sem item".

Regras da auditoria (cada apontamento traz a base legal):
  ICMS-ST      RICMS/RO (Dec. 22.721/2018), Anexo VI, Parte 2 — lista vigente em st_ro_dados.py (sem revogados nem
               redações anteriores; Tabela XXVI só vale para venda porta a porta, por isso não entra na conferência geral).
  Alíquota     Lei 688/96, art. 27, I (internas, só itens vigentes) e art. 27-A (Fecoep +2%); interestaduais pelas
               Resoluções do Senado 22/89 (7%/12%) e 13/12 (4% importados).
  Crédito      LC 87/96, art. 33, I (uso e consumo só a partir de 2033) e art. 20, § 5º (ativo pelo CIAP).
  DIFAL        Lei 688/96, art. 17, XIII e art. 18 (entrada interestadual para uso/consumo ou ativo).
  Combustível  LC 192/2022; Convênios ICMS 199/2022 (diesel, biodiesel, GLP) e 15/2023 (gasolina, etanol anidro): CST 61/60.
  PIS/COFINS   Leis 10.637/2002 e 10.833/2003, art. 3º, I, "b" (sem crédito na compra de monofásico para revenda);
               produtos monofásicos pela Tabela 4.3.10 da EFD-Contribuições (oportunidades.REGRAS_MONOFASICO).
  Cadastro     Guia Prático da EFD ICMS/IPI, registro 0200, campo TIPO_ITEM × CFOP da entrada.
Situação de cada item: Divergente (regra clara, SPED diferente) · Verificar (depende de algo que o SPED não diz) ·
Confere. Itens "Info" só informam (ex.: produto da ST comprado para uso e consumo).
"""
import os
import re
import sys
import json
import html as _h
from collections import defaultdict, OrderedDict

UF_IBGE = {'11': 'RO', '12': 'AC', '13': 'AM', '14': 'RR', '15': 'PA', '16': 'AP', '17': 'TO', '21': 'MA', '22': 'PI', '23': 'CE',
           '24': 'RN', '25': 'PB', '26': 'PE', '27': 'AL', '28': 'SE', '29': 'BA', '31': 'MG', '32': 'ES', '33': 'RJ', '35': 'SP',
           '41': 'PR', '42': 'SC', '43': 'RS', '50': 'MS', '51': 'MT', '52': 'GO', '53': 'DF'}
SUL_SUDESTE_SEM_ES = {'SP', 'RJ', 'MG', 'PR', 'SC', 'RS'}
ORIGEM_IMPORTADA = {'1', '2', '3', '8'}
CANC = {'02', '03', '04', '05'}
CFOP_NOME = {
    '101': 'Compra para industrialização', '102': 'Compra para comercialização', '116': 'Compra para industrialização (entrega futura)',
    '117': 'Compra para comercialização (entrega futura)', '118': 'Compra para comercialização (venda à ordem)',
    '120': 'Compra para comercialização (venda à ordem, já recebida)', '122': 'Compra para industrialização (remetida pelo fornecedor)',
    '124': 'Industrialização efetuada por outra empresa', '126': 'Compra para prestação de serviço', '201': 'Devolução de venda (produção)',
    '202': 'Devolução de venda (revenda)', '251': 'Compra de energia para distribuição', '252': 'Compra de energia (indústria)',
    '253': 'Compra de energia (comércio)', '301': 'Aquisição de comunicação', '351': 'Aquisição de transporte', '352': 'Transporte (indústria)',
    '353': 'Transporte (comércio)', '401': 'Compra para industrialização com ST', '403': 'Compra para comercialização com ST',
    '407': 'Compra de uso e consumo com ST', '406': 'Compra de ativo com ST', '411': 'Devolução de venda com ST',
    '551': 'Compra para o ativo imobilizado', '552': 'Transferência de ativo', '553': 'Devolução de venda de ativo',
    '556': 'Compra de material de uso e consumo', '557': 'Transferência de material de uso e consumo', '651': 'Compra de combustível para industrialização',
    '652': 'Compra de combustível para comercialização', '653': 'Compra de combustível para consumo', '901': 'Entrada para industrialização por encomenda',
    '902': 'Retorno de industrialização por encomenda', '908': 'Entrada de bem em comodato ou locação', '909': 'Retorno de bem em comodato ou locação',
    '910': 'Entrada de bonificação, doação ou brinde', '911': 'Entrada de amostra grátis', '912': 'Entrada para demonstração',
    '913': 'Retorno de demonstração', '915': 'Entrada para conserto', '916': 'Retorno de conserto', '920': 'Entrada de vasilhame',
    '921': 'Retorno de vasilhame', '922': 'Simples faturamento (entrega futura)', '923': 'Entrada por conta e ordem (venda à ordem)',
    '949': 'Outra entrada não especificada', '152': 'Transferência para comercialização', '151': 'Transferência para industrialização'}
CST_NOME = {'00': 'tributada integralmente', '10': 'tributada com cobrança de ST', '20': 'com redução de base', '30': 'isenta/não tributada com ST',
            '40': 'isenta', '41': 'não tributada', '50': 'suspensão', '51': 'diferimento', '60': 'ICMS cobrado anteriormente por ST',
            '61': 'monofásico sobre combustíveis', '70': 'redução de base com ST', '90': 'outras', '02': 'monofásico próprio (combustível)',
            '15': 'monofásico próprio e com retenção', '53': 'monofásico com recolhimento diferido'}
TIPO_ITEM = {'00': 'mercadoria para revenda', '01': 'matéria-prima', '02': 'embalagem', '03': 'produto em processo', '04': 'produto acabado',
             '05': 'subproduto', '06': 'produto intermediário', '07': 'material de uso e consumo', '08': 'ativo imobilizado', '09': 'serviços',
             '10': 'outros insumos', '99': 'outras'}

# Lei 688/96, art. 27, I — alíquotas internas VIGENTES por produto (itens revogados fora). (prefixos NCM, alíquota, rótulo, dispositivo)
# A lei descreve o produto por nome; o NCM é a correspondência usada pelo sistema (conferir em caso de dúvida).
ALIQ_INTERNA = [
    (['7108'], 9.0, 'ouro', 'art. 27, I, "a"'),
    (['7102', '7103', '7104'], 9.0, 'pedras preciosas', 'art. 27, I, "a"'),
    (['01'], 12.0, 'animais vivos', 'art. 27, I, "b", 1'),
    (['0201', '0202', '0203', '0204', '0206', '0207', '020810'], 12.0, 'carnes e miúdos comestíveis', 'art. 27, I, "b", 2'),
    (['0302', '0303'], 12.0, 'peixes frescos, resfriados ou congelados', 'art. 27, I, "b", 3'),
    (['071333'], 12.0, 'feijão', 'art. 27, I, "b", 5'),
    (['110620'], 12.0, 'farinha de mandioca', 'art. 27, I, "b", 6'),
    (['25010020'], 12.0, 'sal de cozinha', 'art. 27, I, "b", 7'),
    (['0407'], 12.0, 'ovos (hortifrutigranjeiros em estado natural)', 'art. 27, I, "b", 8'),
    (['07', '08'], 12.0, 'hortifrutigranjeiros em estado natural', 'art. 27, I, "b", 8'),
    (['1507'], 12.0, 'óleo de soja para consumo humano', 'art. 27, I, "b", 10'),
    (['170114', '170199'], 12.0, 'açúcar cristal', 'art. 27, I, "b", 11'),
    (['1101'], 12.0, 'farinha de trigo', 'art. 27, I, "b", 12'),
    (['0401'], 12.0, 'leite fresco ou pasteurizado (exceto UHT)', 'art. 27, I, "b", 13'),
    (['110220'], 12.0, 'fubá de milho', 'art. 27, I, "b", 14'),
    (['93'], 27.0, 'armas e munições (25% + 2% Fecoep)', 'art. 27, I, "d", 1, e art. 27-A'),
    (['3303', '3304'], 25.0, 'perfumes e cosméticos', 'art. 27, I, "d", 3'),
    (['8903'], 27.0, 'embarcações de esporte e recreação (25% + 2% Fecoep)', 'art. 27, I, "d", 5, e art. 27-A'),
    (['360410'], 27.0, 'fogos de artifício (25% + 2% Fecoep)', 'art. 27, I, "d", 9, e art. 27-A'),
    (['2402', '2403'], 39.0, 'cigarros, charutos e fumo (37% + 2% Fecoep)', 'art. 27, I, "g", e art. 27-A'),
    (['2204', '2205', '2206', '2208'], 39.0, 'bebidas alcoólicas, exceto cerveja (37% + 2% Fecoep)', 'art. 27, I, "h", e art. 27-A'),
    (['2203'], 36.0, 'cerveja, exceto sem álcool (34% + 2% Fecoep)', 'art. 27, I, "k", e art. 27-A'),
]
ALIQ_VERIFICAR = {'0401': 'leite UHT tem o mesmo NCM do leite pasteurizado e vai a 19,5%', '2203': 'cerveja sem álcool não entra nos 34%',
                  '1507': 'só o óleo de soja para consumo humano tem 12%', '07': 'os 12% são só para o produto em estado natural',
                  '08': 'os 12% são só para o produto em estado natural'}
COMB_MONO = ('27101921', '27101259', '271112', '271113', '271119', '3826', '22071010', '22072011')   # diesel, gasolina C, GLP, biodiesel, etanol anidro
COMB_FORTE = ('27101921', '27101259')
FIM_REVENDA = {'102', '403', '101', '401', '117', '118', '120', '122', '652', '651'}
FIM_USO = {'556', '407', '551', '406', '653'}
FIM_ATIVO = {'551', '406'}
COD_DIFAL = {'RO40000001', 'RO40000002', 'RO90000002'}


def _n(s):
    s = (s or '').strip()
    if not s:
        return 0.0
    try:
        return float(s.replace('.', '').replace(',', '.')) if ',' in s else float(s)
    except ValueError:
        return 0.0


def _fm(v):
    return f'{v:,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')


def _pc(v):
    return (f'{v:.2f}'.rstrip('0').rstrip('.') + '%').replace('.', ',')


# ------------------------------------------------------------------ tabelas
def _tabela_st():
    import st_ro_dados as s
    vig = {c: dict(r) for c, r in s.VIGENTES.items()}
    for c, r in vig.items():          # subitem sem NCM no texto herda o NCM do item principal (xx.yyy.00)
        if not r['ncm']:
            pai = vig.get(c[:7] + '00')
            if pai and pai['ncm']:
                r['ncm'] = list(pai['ncm']); r['herdado'] = True
    lista = [r for r in vig.values() if r['tab'] != 'XXVI' and r['ncm']]
    rev = [r for r in s.REVOGADOS if r['ncm']]
    return lista, rev, s


def _melhor(ncm, lista):
    best, tam = None, 0
    for r in lista:
        for p in r['ncm']:
            if ncm.startswith(p) and len(p) > tam:
                best, tam = r, len(p)
    return best


def _aliq_interna(ncm, dt_ini):
    import difal_ro as dr
    geral = dr.aliq_interna_padrao(dt_ini)
    best = None
    for pre, a, rot, disp in ALIQ_INTERNA:
        for p in pre:
            if ncm.startswith(p) and (best is None or len(p) > best[0]):
                best = (len(p), a, rot, disp, p)
    if best:
        return best[1], best[2], best[3], ALIQ_VERIFICAR.get(best[4][:4]) or ALIQ_VERIFICAR.get(best[4][:2] if len(best[4]) == 2 else '')
    return geral, 'demais produtos', 'art. 27, I, "c"', None


# ------------------------------------------------------------------ leitura do SPED
def ler_entradas(texto):
    linhas = texto.split('\n')
    cut = next((i for i, l in enumerate(linhas) if l.startswith('|9999|')), None)
    if cut is not None:
        linhas = linhas[:cut + 1]
    emp, p0200, p0150, p0190 = {}, {}, {}, {}
    itens, notas = [], []
    nota = None

    def fecha():
        if nota and not nota['tem_item']:
            for c190 in nota['c190']:
                itens.append(dict(nota_idx=nota['idx'], cod='(nota sem item)', desc='Nota escriturada só com totais (C190)', ncm='', tipo='',
                                  unid='', qtd=0.0, cfop=c190['cfop'], cst=c190['cst'], vl=c190['vl'], bc=c190['bc'], aliq=c190['aliq'],
                                  icms=c190['icms'], bcst=c190['bcst'], st=c190['st'], ipi=c190['ipi'], cstpis='', vpis=0.0, vcof=0.0,
                                  nat='', num_item='', sem_item=True))

    for l in linhas:
        c = l.rstrip('\r').split('|')
        if len(c) < 3:
            continue
        r = c[1]
        if r == '0000':
            emp = {'nome': c[6] if len(c) > 6 else '', 'cnpj': c[7] if len(c) > 7 else '', 'dt_ini': c[4], 'dt_fin': c[5],
                   'uf': c[9] if len(c) > 9 else ''}
        elif r == '0150':
            p0150[c[2]] = {'nome': c[3], 'cnpj': c[5] or c[6], 'mun': c[8] if len(c) > 8 else ''}
        elif r == '0190':
            p0190[c[2]] = c[3]
        elif r == '0200':
            p0200[c[2]] = {'desc': c[3], 'unid': c[6] if len(c) > 6 else '', 'tipo': c[7] if len(c) > 7 else '',
                           'ncm': re.sub(r'\D', '', c[8] if len(c) > 8 else ''), 'cest': c[13] if len(c) > 13 else ''}
        elif r == 'C100':
            fecha()
            nota = None
            if c[2] != '0' or c[6] in CANC:
                continue
            part = p0150.get(c[4], {})
            ch = c[9] or ''
            uf = UF_IBGE.get(ch[:2], '') if len(ch) == 44 else UF_IBGE.get((part.get('mun') or '')[:2], '')
            nota = {'idx': len(notas), 'num': c[8], 'ser': c[7], 'mod': c[5], 'chave': ch, 'dt': c[10], 'dt_es': c[11],
                    'forn': part.get('nome', ''), 'cnpj': part.get('cnpj', ''), 'uf': uf, 'vl': _n(c[12]), 'tem_item': False,
                    'c190': [], 'c197': []}
            notas.append(nota)
        elif nota is None:
            continue
        elif r == 'C170':
            nota['tem_item'] = True
            p = p0200.get(c[3], {})
            un = c[6] or p.get('unid', '')
            itens.append(dict(nota_idx=nota['idx'], cod=c[3], desc=p.get('desc') or c[4], ncm=p.get('ncm', ''), tipo=p.get('tipo', ''),
                              unid=p0190.get(un, un), qtd=_n(c[5]), cfop=c[11], cst=c[10], vl=_n(c[7]) - _n(c[8]), bc=_n(c[13]),
                              aliq=_n(c[14]), icms=_n(c[15]), bcst=_n(c[16]), st=_n(c[18]), ipi=_n(c[24]) if len(c) > 24 else 0.0,
                              cstpis=c[25] if len(c) > 25 else '', vpis=_n(c[30]) if len(c) > 30 else 0.0,
                              vcof=_n(c[36]) if len(c) > 36 else 0.0, nat=c[12], num_item=c[2], sem_item=False))
        elif r == 'C190':
            nota['c190'].append({'cst': c[2], 'cfop': c[3], 'aliq': _n(c[4]), 'vl': _n(c[5]), 'bc': _n(c[6]), 'icms': _n(c[7]),
                                 'bcst': _n(c[8]), 'st': _n(c[9]), 'ipi': _n(c[11]) if len(c) > 11 else 0.0})
        elif r == 'C197':
            nota['c197'].append(c[2])
    fecha()
    return emp, notas, itens


# ------------------------------------------------------------------ auditoria de um item
def _auditar(i, nota, dt_ini, st_lista, st_rev):
    ach = []
    ncm, cf, cst = i['ncm'], i['cfop'] or '', (i['cst'] or '').zfill(3)
    fim, cst2, orig = cf[1:], cst[-2:], cst[:1]
    if i['sem_item']:
        return [('Verificar', 'Cadastro', 'Nota sem itens no SPED (só C190): não dá para auditar produto a produto',
                 'Guia Prático da EFD ICMS/IPI, registro C170')]
    revenda, uso, ativo = fim in FIM_REVENDA, fim in FIM_USO, fim in FIM_ATIVO
    m = _melhor(ncm, st_lista) if ncm else None
    mr = None if m else (_melhor(ncm, st_rev) if ncm else None)
    tem_st = cst2 in ('10', '30', '60', '70') or fim in ('401', '403', '407', '406', '411') or i['st'] > 0
    # ICMS-ST
    if m:
        tab = f"Tabela {m['tab']}, CEST {m['cest']}"
        if revenda and not tem_st and not ncm.startswith(COMB_MONO):
            ach.append(('Verificar', 'ICMS-ST', f'Produto na ST de RO ({tab}) entrou para revenda sem ST na nota (CST {cst}). '
                        'Conferir se o ICMS-ST foi recolhido na entrada (DARE 1231) ou se há regime especial.', 'RICMS/RO, Anexo VI, Parte 2'))
        elif uso:
            ach.append(('Info', 'ICMS-ST', f'Produto da ST de RO ({tab}) comprado para uso e consumo/ativo', 'RICMS/RO, Anexo VI, Parte 2'))
        else:
            ach.append(('Info', 'ICMS-ST', f'Produto na ST de RO ({tab})', 'RICMS/RO, Anexo VI, Parte 2'))
    elif revenda and tem_st and not ncm.startswith(COMB_MONO):
        if mr:
            quando = f" em {mr['revogado_em']}" if mr.get('revogado_em') else ''
            ach.append(('Divergente', 'ICMS-ST', f"Produto saiu da ST de RO{quando} (Tabela {mr['tab']}, CEST {mr['cest']}, revogado) e entrou com ST "
                        f'(CST {cst} / CFOP {cf})', 'RICMS/RO, Anexo VI, Parte 2'))
        elif ncm:
            ach.append(('Verificar', 'ICMS-ST', f'Entrou com ST (CST {cst} / CFOP {cf}), mas o NCM {ncm} não está na lista de ST de RO',
                        'RICMS/RO, Anexo VI, Parte 2'))
    uso_consumo = fim in FIM_USO and not ativo
    # material de uso e consumo: a entrada é escriturada sem base, alíquota e ICMS (não gera crédito até 2033)
    if uso_consumo and (i['aliq'] > 0 or i['bc'] > 0 or i['icms'] > 0):
        veio = ', '.join(x for x in ((f"base R$ {_fm(i['bc'])}" if i['bc'] > 0 else ''), (f"alíquota {_pc(i['aliq'])}" if i['aliq'] > 0 else ''),
                                     (f"ICMS R$ {_fm(i['icms'])}" if i['icms'] > 0 else '')) if x)
        ach.append(('Divergente', 'Crédito ICMS', f'Mercadoria de uso e consumo (CFOP {cf}) não pode ter destaque de ICMS na escrituração da '
                    f'entrada: base, alíquota e ICMS ficam zerados, porque não há crédito (veio {veio}).',
                    'LC 87/96, art. 33, I; Guia Prático da EFD ICMS/IPI, registro C170'))
    # alíquota (uso e consumo fica de fora: já apontado acima)
    if i['aliq'] > 0 and not uso_consumo:
        if cf[:1] == '1' and cst2 not in ('20', '70'):     # CST 20/70: redução de base (benefício) — a alíquota nominal pode diferir da carga
            a, rot, disp, ver = _aliq_interna(ncm, dt_ini)
            if abs(i['aliq'] - a) > 0.009:
                if i['aliq'] > a and not ver:
                    ach.append(('Divergente', 'Alíquota', f"Compra dentro de RO com alíquota de {_pc(i['aliq'])}, maior que a interna de {_pc(a)} para {rot}: "
                                'o crédito fica limitado à alíquota correta', f'Lei 688/96, {disp}'))
                else:
                    extra = f' ({ver})' if ver else (' Pode haver benefício (isenção, redução ou crédito presumido) que esta versão ainda não confere.'
                                                     if i['aliq'] < a else '')
                    ach.append(('Verificar', 'Alíquota', f"Compra dentro de RO com alíquota de {_pc(i['aliq'])}; para {rot} a alíquota interna é {_pc(a)}.{extra}",
                                f'Lei 688/96, {disp}'))
        elif cf[:1] == '2':
            if orig in ORIGEM_IMPORTADA:
                if abs(i['aliq'] - 4) > 0.009:
                    ach.append(('Divergente', 'Alíquota', f"Produto importado (origem {orig}) de outro estado com {_pc(i['aliq'])}: o correto é 4%",
                                'Resolução do Senado 13/2012'))
            else:
                uf = nota['uf']
                esp = 7.0 if uf in SUL_SUDESTE_SEM_ES else 12.0 if uf else None
                if i['aliq'] not in (4.0, 7.0, 12.0):
                    ach.append(('Divergente', 'Alíquota', f"Compra de outro estado com alíquota de {_pc(i['aliq'])} (as interestaduais são 4%, 7% ou 12%)",
                                'Resoluções do Senado 22/1989 e 13/2012'))
                elif esp and abs(i['aliq'] - esp) > 0.009 and i['aliq'] != 4.0:
                    ach.append(('Divergente', 'Alíquota', f"Fornecedor de {uf} com {_pc(i['aliq'])}: para produto nacional vindo de {uf} a alíquota é {_pc(esp)}",
                                'Resolução do Senado 22/1989'))
                elif i['aliq'] == 4.0:
                    ach.append(('Verificar', 'Alíquota', f'4% com origem {orig} (nacional): os 4% são para importados ou com conteúdo de importação acima de 40%',
                                'Resolução do Senado 13/2012'))
    # conta do ICMS
    if not uso_consumo and i['bc'] > 0 and i['aliq'] > 0 and abs(round(i['bc'] * i['aliq'] / 100, 2) - i['icms']) > 0.05:
        ach.append(('Divergente', 'Crédito ICMS', f"ICMS lançado ({_fm(i['icms'])}) ≠ base × alíquota ({_fm(i['bc'] * i['aliq'] / 100)})",
                    'Guia Prático da EFD ICMS/IPI, registro C170'))
    # crédito em uso e consumo / ativo
    if ativo and i['icms'] > 0:
        ach.append(('Verificar', 'Crédito ICMS', 'Crédito de ICMS em compra de ativo: só é permitido pelo CIAP (1/48 por mês); conferir o bloco G',
                    'LC 87/96, art. 20, § 5º'))
    if cst2 in ('40', '41', '50', '60', '61') and i['icms'] > 0 and not uso_consumo:
        ach.append(('Divergente', 'Crédito ICMS', f"CST {cst} ({CST_NOME.get(cst2, '')}) com ICMS creditado de {_fm(i['icms'])}",
                    'Guia Prático da EFD ICMS/IPI; LC 87/96, art. 20'))
    # DIFAL
    if cf[:1] == '2' and fim in ('556', '407', '551', '406') and not (set(nota['c197']) & COD_DIFAL):
        ach.append(('Verificar', 'DIFAL', 'Compra de outro estado para uso e consumo / ativo sem o ajuste do DIFAL na nota (C197 RO40000001/2). '
                    'Conferir se o DIFAL foi lançado de outra forma ou pago em DARE (1660/1661).', 'Lei 688/96, art. 17, XIII, e art. 18'))
    # combustível monofásico
    if ncm.startswith(COMB_MONO) and cst2 not in ('61', '60', '02', '15', '53'):
        ach.append(('Divergente' if ncm.startswith(COMB_FORTE) else 'Verificar', 'Combustível',
                    f'Combustível com CST {cst}: no ICMS monofásico a entrada usa CST 61 (ou 60)',
                    'LC 192/2022; Convênios ICMS 199/2022 e 15/2023'))
    # PIS/COFINS monofásico
    try:
        from oportunidades import _regra_ncm
        rg = _regra_ncm(ncm) if ncm else None
    except Exception:  # pragma: no cover
        rg = None
    if rg:
        if revenda and (i['cstpis'] or '') in ('50', '51', '52', '53', '54', '55', '56') and (i['vpis'] > 0 or i['vcof'] > 0):
            ach.append(('Divergente', 'PIS/COFINS', f"Crédito de PIS/COFINS na compra para revenda de produto monofásico ({rg['grupo']})",
                        'Lei 10.637/2002 e Lei 10.833/2003, art. 3º, I, "b"'))
        else:
            ach.append(('Info', 'PIS/COFINS', f"Monofásico de PIS/COFINS: {rg['grupo']}", rg['lei']))
    # cadastro × CFOP
    t = i['tipo']
    if t == '00' and uso:
        ach.append(('Verificar', 'Cadastro', f'Produto cadastrado como mercadoria para revenda (0200 tipo 00) entrando como uso e consumo/ativo (CFOP {cf})',
                    'Guia Prático da EFD ICMS/IPI, registro 0200'))
    elif t in ('07', '08') and revenda:
        ach.append(('Verificar', 'Cadastro', f"Produto cadastrado como {TIPO_ITEM[t]} (0200 tipo {t}) entrando para revenda/industrialização (CFOP {cf})",
                    'Guia Prático da EFD ICMS/IPI, registro 0200'))
    return ach


def _acum7000(notas, itens):
    """Itens no acumulador genérico 7000 (COD_NAT do C170), agrupados por nota e, dentro da nota, pelo acumulador de despesa
    sugerido (acumulador_despesa.classifica: descrição + NCM + CFOP). Nota com mais de um acumulador = MISTA (só destaca)."""
    try:
        import acumulador_despesa as ad
    except Exception:  # pragma: no cover
        return {'notas': [], 'nomes': {}}
    por_nota = OrderedDict()
    for i in itens:
        if i['sem_item'] or (i.get('nat') or '').strip() != ad.ACUM_GENERICO:
            continue
        sug, motivo, conf = ad.classifica(i['desc'], i['ncm'], i['cfop'])
        por_nota.setdefault(i['nota_idx'], []).append({'item': i['num_item'], 'cod': i['cod'], 'desc': i['desc'], 'ncm': i['ncm'],
                                                     'cfop': i['cfop'], 'vl': i['vl'], 'acum': sug or '', 'motivo': motivo, 'conf': conf})
    out = []
    for ni, its in por_nota.items():
        n = notas[ni]
        grupos = OrderedDict()
        for it in sorted(its, key=lambda x: -x['vl']):
            grupos.setdefault(it['acum'], []).append(it)
        partes = sorted(([a, g] for a, g in grupos.items()), key=lambda x: -sum(y['vl'] for y in x[1]))
        out.append({'num': n['num'], 'dt': n['dt_es'] or n['dt'], 'forn': n['forn'], 'cnpj': n['cnpj'], 'vl': sum(x['vl'] for x in its),
                    'mista': len(partes) > 1, 'partes': [{'acum': a, 'vl': sum(y['vl'] for y in g), 'itens': g,
                                                          'conf': 'media' if any(y['conf'] != 'alta' for y in g) else 'alta'} for a, g in partes]})
    return {'notas': out, 'nomes': dict(ad.ACUM)}


# Tipo de operação da entrada pelo CFOP (3 últimos dígitos) — Ajuste SINIEF s/nº/1970 (Tabela de CFOP)
TIPO_OPER = [
    ('Revenda', {'102', '403', '117', '118', '120', '652', '113', '114', '115', '121'}),
    ('Industrialização / insumo', {'101', '401', '116', '122', '124', '125', '651', '111', '126', '128'}),
    ('Uso e consumo', {'556', '407', '653', '557'}),
    ('Ativo imobilizado', {'551', '406', '552', '553', '554', '555'}),
    ('Devolução de venda', {'201', '202', '203', '204', '208', '209', '410', '411', '412', '413', '503', '504', '660', '661', '662'}),
    ('Transferência', {'151', '152', '153', '154', '408', '409', '658', '659'}),
    ('Energia, comunicação e transporte', {'251', '252', '253', '254', '255', '256', '257', '301', '302', '303', '304', '305', '306',
                                           '351', '352', '353', '354', '355', '356', '360'}),
    ('Bonificação, brinde e amostra', {'910', '911'}),
    ('Remessas e retornos (conserto, demonstração, industrialização, comodato, vasilhame)',
     {'901', '902', '903', '904', '905', '906', '907', '908', '909', '912', '913', '914', '915', '916', '917', '918', '919', '920', '921', '924',
      '925', '934'}),
    ('Venda à ordem / entrega futura (simbólicas)', {'922', '923'}),
]


def tipo_operacao(cfop):
    fim = (cfop or '')[1:]
    for nome, finais in TIPO_OPER:
        if fim in finais:
            return nome
    return 'Outras entradas'


ORDEM = {'Confere': 0, 'Info': 0, 'Verificar': 1, 'Divergente': 2}
TEMAS = [('ICMS-ST', 'Produto na ST de RO × ST da nota'), ('Alíquota', 'Interna (Lei 688, art. 27) ou interestadual'),
         ('Crédito ICMS', 'Crédito em uso e consumo, CST e conta do ICMS'), ('DIFAL', 'Entrada interestadual para uso e consumo / ativo'),
         ('Combustível', 'ICMS monofásico: CST 61 / 60'), ('PIS/COFINS', 'Monofásico: crédito na compra'),
         ('Cadastro', 'Tipo do produto no 0200 × CFOP')]


def analisar_texto(texto):
    import difal_ro as dr  # noqa: F401  (alíquota geral por período)
    emp, notas, itens = ler_entradas(texto)
    st_lista, st_rev, st_mod = _tabela_st()
    for i in itens:
        i['ach'] = _auditar(i, notas[i['nota_idx']], emp.get('dt_ini', ''), st_lista, st_rev)
        i['sit'] = max((a[0] for a in i['ach']), key=lambda s: ORDEM[s], default='Confere')
        if i['sit'] == 'Info':
            i['sit'] = 'Confere'
    prod = OrderedDict()
    for k, i in enumerate(itens):
        key = i['cod'] if not i['sem_item'] else f"(nota sem item) {notas[i['nota_idx']]['num']}"
        g = prod.get(key)
        if not g:
            g = prod[key] = {'cod': i['cod'], 'desc': i['desc'], 'ncm': i['ncm'], 'tipo': i['tipo'], 'itens': [], 'vl': 0.0, 'bc': 0.0,
                             'icms': 0.0, 'st': 0.0, 'ipi': 0.0, 'qtd': 0.0, 'sit': 'Confere', 'ach': OrderedDict()}
        g['itens'].append(k)
        for f in ('vl', 'bc', 'icms', 'st', 'ipi', 'qtd'):
            g[f] += i[f]
        if ORDEM[i['sit']] > ORDEM[g['sit']]:
            g['sit'] = i['sit']
        for a in i['ach']:
            g['ach'].setdefault((a[0], a[1], a[2]), a[3])
    for g in prod.values():
        g['cfops'] = sorted({itens[k]['cfop'] for k in g['itens']})
        g['csts'] = sorted({itens[k]['cst'] for k in g['itens']})
        g['aliqs'] = sorted({itens[k]['aliq'] for k in g['itens']})
        g['notas'] = len({itens[k]['nota_idx'] for k in g['itens']})
        g['ach'] = [list(k) + [v] for k, v in g['ach'].items()]

    def grupo(chave):
        d = defaultdict(lambda: {'n': 0, 'vl': 0.0, 'bc': 0.0, 'icms': 0.0, 'st': 0.0})
        for i in itens:
            x = d[chave(i)]
            x['n'] += 1
            for f in ('vl', 'bc', 'icms', 'st'):
                x[f] += i[f]
        return sorted(([k, v] for k, v in d.items()), key=lambda x: -x[1]['vl'])

    tot = {f: sum(i[f] for i in itens) for f in ('vl', 'bc', 'icms', 'st', 'ipi')}
    acum7000 = _acum7000(notas, itens)
    sit_p = defaultdict(lambda: [0, 0.0])
    for g in prod.values():
        sit_p[g['sit']][0] += 1; sit_p[g['sit']][1] += g['vl']
    temas = {}
    for t, _ in TEMAS:
        temas[t] = {s: sum(1 for g in prod.values() if any(a[1] == t and a[0] == s for a in g['ach'])) for s in ('Divergente', 'Verificar', 'Info')}
    pontos = []
    v2 = sum(i['vl'] for i in itens if i['cfop'][:1] == '2' and i['cfop'][1:] in ('556', '407', '551', '406'))
    if v2:
        pontos.append(['a', f'R$ {_fm(v2)} em uso e consumo / ativo de outro estado (2556, 2407, 2551, 2406): confira o DIFAL.'])
    ncr = sum(1 for i in itens if i['icms'] > 0)
    pontos.append(['i', f'{ncr} item(ns) com crédito de ICMS, somando R$ {_fm(tot["icms"])}.'])
    nsem = sum(1 for n in notas if not n['tem_item'])
    if nsem:
        pontos.append(['a', f'{nsem} nota(s) de entrada sem itens no SPED (só C190): entram nos totais, mas não na auditoria por produto.'])
    return {
        'empresa': emp, 'notas': [{k: n[k] for k in ('num', 'ser', 'mod', 'chave', 'dt', 'dt_es', 'forn', 'cnpj', 'uf', 'vl', 'tem_item')} for n in notas],
        'itens': [{k: v for k, v in i.items()} for i in itens], 'produtos': list(prod.values()), 'totais': tot,
        'por_cfop': grupo(lambda i: i['cfop']), 'por_cst': grupo(lambda i: i['cst']), 'por_aliq': grupo(lambda i: i['aliq']),
        'por_forn': grupo(lambda i: notas[i['nota_idx']]['forn'] or '(sem nome)'),
        'por_tipo': grupo(lambda i: tipo_operacao(i['cfop'])),
        'situacao': {k: v for k, v in sit_p.items()}, 'temas': temas, 'pontos': pontos, 'acum7000': acum7000,
        'fontes': {'st': f'RICMS/RO, Anexo VI, Parte 2 — consolidado até o {st_mod.VERSAO}: {len(st_mod.VIGENTES)} CEST vigentes, '
                         f'{len(st_mod.REVOGADOS)} revogados descartados; Tabela XXVI (porta a porta) fora da conferência geral',
                   'aliq': 'Lei 688/96, art. 27, I, e art. 27-A (Fecoep), sem os itens revogados',
                   'pis': 'Tabela 4.3.10 da EFD-Contribuições (monofásicos)',
                   'pendente': 'Benefícios de RO (Anexos I e II do RICMS: isenção e redução de base) ainda não entram nesta versão'},
    }


# ------------------------------------------------------------------ HTML (animado, interativo, autocontido)
_CSS = r'''
:root{--bg:#f4f6fb;--pn:#fff;--ink:#1b2333;--ink2:#64748b;--ln:#e3e7ef;--nv:#1e3a8a;--or:#ea580c;--ok:#0ea472;--dv:#e23d4c;--vf:#f59e0b;--in:#6366f1}
*{box-sizing:border-box}body{margin:0;background:var(--bg);font-family:Inter,"Segoe UI",Arial,sans-serif;color:var(--ink);font-size:13px}
.top{background:linear-gradient(120deg,#0f1b45,#1e3a8a 70%,#2b3a72);color:#fff;padding:20px 28px 0}
.top small{opacity:.75;letter-spacing:.12em;font-size:11px;font-weight:700}.top h1{margin:6px 0 4px;font-size:22px}.top p{margin:0 0 16px;opacity:.85}
.tabs{display:flex;gap:6px}.tabs button{border:0;cursor:pointer;padding:9px 18px;border-radius:9px 9px 0 0;font:inherit;font-weight:700;background:rgba(255,255,255,.12);color:#fff}
.tabs button.on{background:var(--bg);color:var(--nv)}
.wrap{padding:16px 28px 30px}.pg{display:none}.pg.on{display:block;animation:sobe .4s ease both}
@keyframes sobe{from{opacity:0;transform:translateY(10px)}to{opacity:1;transform:none}}
.k{display:grid;grid-template-columns:repeat(auto-fit,minmax(170px,1fr));gap:12px}.k>div{background:var(--pn);border-radius:12px;padding:13px 15px;box-shadow:0 2px 10px rgba(15,23,42,.06);border-left:5px solid var(--nv);animation:sobe .5s ease both}
.k span{font-size:10.5px;color:var(--ink2);text-transform:uppercase;letter-spacing:.05em}.k b{display:block;font-size:21px;margin-top:3px}.k small{font-size:11px;color:var(--ink2)}
.g2{display:grid;grid-template-columns:1.2fr 1fr;gap:14px;margin-top:14px}.c{background:var(--pn);border-radius:12px;padding:15px 17px;box-shadow:0 2px 10px rgba(15,23,42,.05);margin-top:14px}.g2 .c{margin-top:0}
.c h3{margin:0 0 10px;font-size:14px}
.br{display:grid;grid-template-columns:minmax(120px,240px) 1fr 120px;gap:3px 10px;align-items:center;font-size:12px;margin:7px 0}
.br i{height:11px;background:#eef2f7;border-radius:6px;overflow:hidden}.br em{display:block;height:100%;width:0;background:linear-gradient(90deg,#1e3a8a,#3b82f6);border-radius:6px;transition:width 1s ease}
.br b{text-align:right;font-variant-numeric:tabular-nums}.br small{grid-column:2/4;color:var(--ink2);font-size:11px;margin-top:-3px}
.al{border-radius:10px;padding:9px 12px;font-size:12.5px;margin:7px 0;border-left:4px solid}.al.a{background:#fff7ed;border-color:var(--or);color:#7c2d12}.al.i{background:#eff6ff;border-color:var(--nv);color:var(--nv)}
.f{display:flex;gap:7px;align-items:center;margin-bottom:10px;flex-wrap:wrap}.f span{font-size:12px;border:1px solid #d6dbe6;border-radius:999px;padding:5px 11px;font-weight:600;cursor:pointer;user-select:none}
.f span.on{background:var(--nv);color:#fff;border-color:var(--nv)}.f input,.f select{border:1px solid #d6dbe6;border-radius:8px;padding:7px 10px;font:inherit;font-size:12px;background:#fff}.f input{margin-left:auto;width:270px}
table{width:100%;border-collapse:collapse;font-size:12px}th{text-align:left;font-size:10.5px;text-transform:uppercase;color:var(--ink2);padding:8px;border-bottom:1px solid var(--ln);background:#f8fafc;position:sticky;top:0}
td{padding:7px 8px;border-bottom:1px solid #eef1f6;vertical-align:top}td small{display:block;color:var(--ink2);font-size:11px;margin-top:2px}.r{text-align:right;font-variant-numeric:tabular-nums;white-space:nowrap}
tr.pr{cursor:pointer}tr.pr:hover td{background:#f8fafc}tr.det td{background:#fbfcfe;padding:4px 8px 12px 28px}.det table{font-size:11.5px}
.s{font-size:11px;font-weight:800;border-radius:999px;padding:3px 10px;white-space:nowrap}.s.Divergente{background:#fee2e2;color:#991b1b}.s.Verificar{background:#fef3c7;color:#92400e}.s.Confere{background:#dcfce7;color:#166534}.s.Info{background:#e0e7ff;color:#3730a3}
.tm{display:grid;grid-template-columns:repeat(auto-fit,minmax(250px,1fr));gap:12px;margin-top:14px}.tm>div{background:var(--pn);border-radius:12px;padding:12px 14px;border:1px solid var(--ln);border-top:4px solid var(--ok);cursor:pointer;animation:sobe .5s ease both}
.tm>div.dv{border-top-color:var(--dv)}.tm>div.vf{border-top-color:var(--vf)}.tm>div.on{outline:2px solid var(--nv)}.tm b{font-size:14px}.tm p{margin:3px 0 8px;font-size:12px;color:#475569}
.b{display:inline-block;white-space:nowrap;font-size:11px;font-weight:700;border-radius:999px;padding:2px 9px;margin:0 5px 4px 0}.b.d{background:#fee2e2;color:#991b1b}.b.v{background:#fef3c7;color:#92400e}.b.n{background:#e0e7ff;color:#3730a3}
.ft{font-size:11.5px;color:var(--ink2);margin-top:14px;background:var(--pn);border:1px dashed #cbd5e1;border-radius:10px;padding:10px 14px;line-height:1.6}
.mais[hidden]{display:none}.mais{display:block;margin:10px auto 0;border:1px solid #c7d2fe;background:#eef2ff;color:var(--nv);font-weight:700;border-radius:8px;padding:7px 14px;cursor:pointer;font:inherit}
.donut{display:flex;gap:16px;align-items:center}.lg{display:grid;grid-template-columns:12px 42px 1fr 48px 110px;gap:7px;align-items:center;font-size:12px;margin:4px 0}.lg i{width:10px;height:10px;border-radius:3px}.lg em{font-style:normal;text-align:right;font-weight:700}.lg u{text-decoration:none;text-align:right;color:#475569}
@media print{.tabs,.f,.mais{display:none}.pg{display:block!important}}
'''

_JS = r'''
const D=JSON.parse(document.getElementById('dados').textContent);
const fm=v=>Number(v||0).toLocaleString('pt-BR',{minimumFractionDigits:2,maximumFractionDigits:2});
const pc=v=>(Math.round(v*100)/100).toLocaleString('pt-BR')+'%';
const es=s=>String(s==null?'':s).replace(/[&<>"]/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;'}[c]));
const $=q=>document.querySelector(q),$$=q=>[...document.querySelectorAll(q)];
const dt=s=>s&&s.length===8?s.slice(0,2)+'/'+s.slice(2,4)+'/'+s.slice(4):s||'';
$$('.tabs button').forEach(b=>b.onclick=()=>{$$('.tabs button').forEach(x=>x.classList.toggle('on',x===b));$$('.pg').forEach(p=>p.classList.toggle('on',p.id===b.dataset.p));anima();});
function anima(){setTimeout(()=>$$('.br em').forEach(e=>e.style.width=e.dataset.w),60);$$('[data-cnt]').forEach(e=>{const alvo=+e.dataset.cnt,din=e.dataset.din==='1';let t0=null;const f=t=>{t0=t0||t;const p=Math.min(1,(t-t0)/700);const v=alvo*(1-Math.pow(1-p,3));e.textContent=din?'R$ '+fm(v):Math.round(v).toLocaleString('pt-BR');if(p<1)requestAnimationFrame(f)};requestAnimationFrame(f)});}
function notasDe(p){return p.itens.map(k=>{const i=D.itens[k],n=D.notas[i.nota_idx];return `<tr><td>${es(n.num)}</td><td>${dt(n.dt_es||n.dt)}</td><td>${es(n.forn)} ${n.uf?'<small>'+es(n.uf)+'</small>':''}</td><td>${es(i.cfop)}</td><td>${es(i.cst)}</td><td class="r">${(i.qtd||0).toLocaleString('pt-BR')} ${es(i.unid)}</td><td class="r">${fm(i.vl)}</td><td class="r">${fm(i.bc)}</td><td class="r">${pc(i.aliq)}</td><td class="r">${fm(i.icms)}</td><td class="r">${fm(i.st)}</td></tr>`}).join('')}
const cabN='<tr><th>Nota</th><th>Entrada</th><th>Fornecedor</th><th>CFOP</th><th>CST</th><th class="r">Qtd</th><th class="r">Valor</th><th class="r">Base</th><th class="r">Alíq.</th><th class="r">ICMS</th><th class="r">ICMS-ST</th></tr>';
// ---------- relatório
const R={f:'todos',q:'',cf:'',cs:'',lim:200};
function filtraR(p){if(R.f==='cred'&&!(p.icms>0))return false;if(R.f==='sem'&&p.icms>0)return false;
 if(R.f==='uso'&&!p.cfops.some(c=>['556','407','551','406','653'].includes(c.slice(1))))return false;
 if(R.f==='rev'&&!p.cfops.some(c=>['102','403','101','401','652'].includes(c.slice(1))))return false;
 if(R.cf&&!p.cfops.includes(R.cf))return false;if(R.cs&&!p.csts.includes(R.cs))return false;
 if(R.q){const t=(p.cod+' '+p.desc+' '+p.ncm+' '+p.itens.map(k=>D.notas[D.itens[k].nota_idx].forn).join(' ')).toLowerCase();if(!t.includes(R.q))return false}return true}
function desenhaR(){const L=D.produtos.filter(filtraR);const tb=$('#tR');
 tb.innerHTML=L.slice(0,R.lim).map((p,j)=>`<tr class="pr" data-j="${D.produtos.indexOf(p)}"><td>${es(p.cod)}</td><td><b>${es(p.desc)}</b></td><td>${es(p.ncm)}</td><td>${p.cfops.join('/')}</td><td>${p.csts.join('/')}</td><td class="r">${fm(p.vl)}</td><td class="r">${fm(p.bc)}</td><td class="r">${p.aliqs.map(pc).join('/')}</td><td class="r">${fm(p.icms)}</td><td class="r">${fm(p.st)}</td><td class="r">${p.notas}</td></tr>`).join('');
 $('#nR').textContent=`Mostrando ${Math.min(L.length,R.lim)} de ${L.length} produtos · clique no produto para ver as notas`;$('#mR').hidden=L.length<=R.lim;liga('#tR',11)}
function liga(sel,cols){$$(sel+' tr.pr').forEach(tr=>tr.onclick=()=>{const nx=tr.nextElementSibling;if(nx&&nx.classList.contains('det')){nx.remove();return}
 const p=D.produtos[+tr.dataset.j];const ach=p.ach.length?`<div style="margin:6px 0 8px">${p.ach.map(a=>`<div style="margin:3px 0"><span class="s ${a[0]}">${a[0]==='Info'?'Informativo':a[0]}</span> <b>${es(a[1])}</b> · ${es(a[2])} <small style="display:inline;color:#1e3a8a">${es(a[3])}</small></div>`).join('')}</div>`:'';
 const d=document.createElement('tr');d.className='det';d.innerHTML=`<td colspan="${cols}">${ach}<table>${cabN}${notasDe(p)}</table></td>`;tr.after(d)})}
$$('#fR span').forEach(s=>s.onclick=()=>{$$('#fR span').forEach(x=>x.classList.toggle('on',x===s));R.f=s.dataset.f;R.lim=200;desenhaR()});
$('#qR').oninput=e=>{R.q=e.target.value.toLowerCase();R.lim=200;desenhaR()};$('#cfR').onchange=e=>{R.cf=e.target.value;desenhaR()};$('#csR').onchange=e=>{R.cs=e.target.value;desenhaR()};
$('#mR').onclick=()=>{R.lim+=300;desenhaR()};
// ---------- auditoria
const A={f:'Divergente',t:'',q:'',lim:200};
function filtraA(p){if(A.f!=='todos'&&p.sit!==A.f)return false;if(A.t&&!p.ach.some(a=>a[1]===A.t&&(A.f==='todos'||a[0]===A.f||A.f==='Confere')))return false;
 if(A.q){const t=(p.cod+' '+p.desc+' '+p.ncm).toLowerCase();if(!t.includes(A.q))return false}return true}
function desenhaA(){const ord={Divergente:2,Verificar:1,Confere:0};const L=D.produtos.filter(filtraA).sort((a,b)=>ord[b.sit]-ord[a.sit]||b.vl-a.vl);
 $('#tA').innerHTML=L.slice(0,A.lim).map(p=>{const pri=p.ach.filter(a=>a[0]!=='Info');const mostra=(A.t?p.ach.filter(a=>a[1]===A.t):(pri.length?pri:p.ach));
 return `<tr class="pr" data-j="${D.produtos.indexOf(p)}"><td><span class="s ${p.sit}">${p.sit}</span></td><td><b>${es(p.desc)}</b><small>cód. ${es(p.cod)} · NCM ${es(p.ncm||'—')} · ${p.notas} nota(s)</small></td><td>${p.cfops.join('/')} · CST ${p.csts.join('/')} · ${p.aliqs.map(pc).join('/')}</td><td>${mostra.map(a=>`<div>${a[0]==='Info'?'<span class="s Info">Informativo</span> ':''}${es(a[2])}<small>${es(a[1])}</small></div>`).join('')||'<small>Sem apontamentos</small>'}</td><td style="font-size:11px;color:#1e3a8a">${[...new Set(mostra.map(a=>a[3]))].map(es).join('<br>')}</td><td class="r">${fm(p.vl)}</td></tr>`}).join('');
 $('#nA').textContent=`Mostrando ${Math.min(L.length,A.lim)} de ${L.length} produtos · clique no produto para ver as notas`;$('#mA').hidden=L.length<=A.lim;liga('#tA',6)}
$$('#fA span').forEach(s=>s.onclick=()=>{$$('#fA span').forEach(x=>x.classList.toggle('on',x===s));A.f=s.dataset.f;A.lim=200;desenhaA()});
$$('.tm>div').forEach(c=>c.onclick=()=>{const on=A.t===c.dataset.t;A.t=on?'':c.dataset.t;$$('.tm>div').forEach(x=>x.classList.toggle('on',!on&&x===c));if(!on){A.f='todos';$$('#fA span').forEach(x=>x.classList.toggle('on',x.dataset.f==='todos'))}desenhaA()});
$('#qA').oninput=e=>{A.q=e.target.value.toLowerCase();A.lim=200;desenhaA()};$('#mA').onclick=()=>{A.lim+=300;desenhaA()};
if(!D.produtos.some(p=>p.sit==='Divergente')){A.f=D.produtos.some(p=>p.sit==='Verificar')?'Verificar':'todos';$$('#fA span').forEach(x=>x.classList.toggle('on',x.dataset.f===A.f))}
desenhaR();desenhaA();anima();
'''


def gerar_html(res):
    e, t = res['empresa'], res['totais']
    cnpj = e.get('cnpj', '')
    cnpj_f = f'{cnpj[:2]}.{cnpj[2:5]}.{cnpj[5:8]}/{cnpj[8:12]}-{cnpj[12:]}' if len(cnpj) == 14 else cnpj
    d = lambda s: f'{s[:2]}/{s[2:4]}/{s[4:]}' if len(s or '') == 8 else (s or '')
    sit = res['situacao']
    nprod, nit, nnot = len(res['produtos']), len(res['itens']), len(res['notas'])
    cores = ['#1e3a8a', '#ea580c', '#0ea472', '#7c3aed', '#0891b2', '#e23d4c', '#ca8a04', '#64748b']
    vt = t['vl'] or 1
    acc, segs, leg = 0.0, '', ''
    cf = res['por_cfop']
    outros = sum(x[1]['vl'] for x in cf[7:])
    top = cf[:7] + ([['outros', {'vl': outros, 'n': 0}]] if outros else [])
    for k, (c, g) in enumerate(top):
        a = g['vl'] / vt
        segs += (f'<circle r="60" cx="80" cy="80" fill="none" stroke="{cores[k % 8]}" stroke-width="26" stroke-dasharray="{a * 377:.2f} 377" '
                 f'stroke-dashoffset="{-acc * 377:.2f}" transform="rotate(-90 80 80)"><animate attributeName="stroke-dasharray" from="0 377" '
                 f'to="{a * 377:.2f} 377" dur="0.9s" fill="freeze"/></circle>')
        acc += a
        nome = CFOP_NOME.get(c[1:], '') if c != 'outros' else 'demais CFOPs'
        leg += (f'<div class="lg"><i style="background:{cores[k % 8]}"></i><b>{_h.escape(c)}</b><span>{_h.escape(nome)}</span>'
                f'<em>{a * 100:.1f}%</em><u>R$ {_fm(g["vl"])}</u></div>')

    def barras(lista, rot, n=8):
        mx = max([g['vl'] for _, g in lista[:n]] or [1]) or 1
        return ''.join(f'<div class="br"><span>{_h.escape(rot(k))}</span><i><em data-w="{max(g["vl"] / mx * 100, 1.2):.1f}%"></em></i>'
                       f'<b>R$ {_fm(g["vl"])}</b><small>{g["n"]} item(ns) · ICMS R$ {_fm(g["icms"])}</small></div>' for k, g in lista[:n])
    b_cst = barras(res['por_cst'], lambda c: f'{c} · {CST_NOME.get((c or "")[-2:], "")}')
    b_al = barras(res['por_aliq'], lambda a: 'sem destaque (0%)' if not a else _pc(a))
    b_fo = barras(res['por_forn'], lambda f: f[:40])
    b_ti = barras(res['por_tipo'], lambda t: t, 12)
    pontos = ''.join(f'<div class="al {c}">{"⚠" if c == "a" else "ⓘ"} {_h.escape(tx)}</div>' for c, tx in res['pontos'])
    op_cf = ''.join(f'<option value="{_h.escape(c)}">{_h.escape(c)}</option>' for c, _ in res['por_cfop'])
    op_cs = ''.join(f'<option value="{_h.escape(c)}">{_h.escape(c)}</option>' for c, _ in res['por_cst'])
    temas = ''
    for tt, desc in TEMAS:
        q = res['temas'][tt]
        cl = 'dv' if q['Divergente'] else 'vf' if q['Verificar'] else ''
        temas += (f'<div class="{cl}" data-t="{_h.escape(tt)}"><b>{tt}</b><p>{desc}</p><span class="b d">{q["Divergente"]} divergente(s)</span>'
                  f'<span class="b v">{q["Verificar"]} verificar</span><span class="b n">{q["Info"]} informativo(s)</span></div>')
    ninfo = sum(1 for p in res['produtos'] if p['sit'] == 'Confere' and any(a[0] == 'Info' for a in p['ach']))
    k = lambda s: sit.get(s, [0, 0.0])
    fontes = res['fontes']
    dados = json.dumps({'produtos': res['produtos'], 'itens': res['itens'], 'notas': res['notas']}, ensure_ascii=False).replace('</', '<\\/')
    return f'''<!doctype html><html lang="pt-BR"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Tributação das entradas · {_h.escape(e.get("nome", ""))}</title><style>{_CSS}</style></head><body>
<div class="top"><small>FISCOCONT+ · ANÁLISE TRIBUTÁRIA · TRIBUTAÇÃO DAS ENTRADAS POR PRODUTO</small><h1>{_h.escape(e.get("nome", ""))}</h1>
<p>CNPJ {cnpj_f} · {d(e.get("dt_ini", ""))} a {d(e.get("dt_fin", ""))} · {nnot} notas de entrada · {nit} itens · {nprod} produtos · Fonte: SPED Fiscal (C100 de entrada, C170, C190, 0200)</p>
<div class="tabs"><button class="on" data-p="pR">Relatório de tributação</button><button data-p="pA">Auditoria de tributação</button></div></div>
<div class="wrap">
<section class="pg on" id="pR">
<div class="k"><div><span>Valor das entradas</span><b data-cnt="{t["vl"]:.2f}" data-din="1">R$ {_fm(t["vl"])}</b><small>{nnot} notas · {nit} itens</small></div>
<div style="border-color:#7c3aed"><span>Produtos diferentes</span><b data-cnt="{nprod}">{nprod}</b><small>pelo código do 0200</small></div>
<div style="border-color:#0891b2"><span>Base do ICMS</span><b data-cnt="{t["bc"]:.2f}" data-din="1">R$ {_fm(t["bc"])}</b><small>{t["bc"] / vt * 100:.2f}% do valor</small></div>
<div style="border-color:var(--ok)"><span>ICMS creditado</span><b data-cnt="{t["icms"]:.2f}" data-din="1">R$ {_fm(t["icms"])}</b><small>{sum(1 for i in res["itens"] if i["icms"] > 0)} itens com crédito</small></div>
<div style="border-color:var(--or)"><span>ICMS-ST</span><b data-cnt="{t["st"]:.2f}" data-din="1">R$ {_fm(t["st"])}</b><small>nos itens</small></div>
<div style="border-color:#64748b"><span>IPI</span><b data-cnt="{t["ipi"]:.2f}" data-din="1">R$ {_fm(t["ipi"])}</b><small>nos itens</small></div></div>
<div class="g2"><div class="c"><h3>Entradas por CFOP</h3><div class="donut"><svg width="160" height="160" viewBox="0 0 160 160">{segs}<text x="80" y="76" text-anchor="middle" font-size="11" fill="#64748b">CFOPs</text><text x="80" y="97" text-anchor="middle" font-size="21" font-weight="800">{len(cf)}</text></svg><div style="flex:1">{leg}</div></div></div>
<div class="c"><h3>Por CST do ICMS</h3>{b_cst}<h3 style="margin-top:14px">Pontos para conferir</h3>{pontos}</div></div>
<div class="c"><h3>Por tipo de operação</h3>{b_ti}</div>
<div class="g2" style="margin-top:14px"><div class="c"><h3>Por alíquota do ICMS</h3>{b_al}</div><div class="c"><h3>Maiores fornecedores</h3>{b_fo}</div></div>
<div class="c"><h3>Produtos · tributação item a item</h3>
<div class="f" id="fR"><span class="on" data-f="todos">Todos</span><span data-f="cred">Com crédito de ICMS</span><span data-f="sem">Sem crédito</span><span data-f="uso">Uso e consumo / ativo</span><span data-f="rev">Revenda / industrialização</span>
<select id="cfR"><option value="">CFOP: todos</option>{op_cf}</select><select id="csR"><option value="">CST: todos</option>{op_cs}</select><input id="qR" placeholder="Buscar código, descrição, NCM ou fornecedor"></div>
<div style="max-height:620px;overflow:auto"><table><thead><tr><th>Código</th><th>Descrição</th><th>NCM</th><th>CFOP</th><th>CST</th><th class="r">Valor</th><th class="r">Base ICMS</th><th class="r">Alíq.</th><th class="r">ICMS</th><th class="r">ICMS-ST</th><th class="r">Notas</th></tr></thead><tbody id="tR"></tbody></table></div>
<small id="nR" style="display:block;text-align:center;color:#64748b;margin-top:8px"></small><button class="mais" id="mR">Mostrar mais</button></div>
</section>
<section class="pg" id="pA">
<div class="k"><div style="border-color:var(--ok)"><span>Confere</span><b style="color:var(--ok)" data-cnt="{k("Confere")[0]}">{k("Confere")[0]}</b><small>produtos · R$ {_fm(k("Confere")[1])}</small></div>
<div style="border-color:var(--dv)"><span>Divergente</span><b style="color:var(--dv)" data-cnt="{k("Divergente")[0]}">{k("Divergente")[0]}</b><small>produtos · R$ {_fm(k("Divergente")[1])}</small></div>
<div style="border-color:var(--vf)"><span>Verificar</span><b style="color:#b45309" data-cnt="{k("Verificar")[0]}">{k("Verificar")[0]}</b><small>produtos · R$ {_fm(k("Verificar")[1])}</small></div>
<div style="border-color:var(--in)"><span>Só informativos</span><b style="color:#4f46e5" data-cnt="{ninfo}">{ninfo}</b><small>conferem, com observação (ex.: produto da ST)</small></div></div>
<div class="tm">{temas}</div>
<div class="c"><h3>Produtos auditados</h3><div class="f" id="fA"><span class="on" data-f="Divergente">Divergente</span><span data-f="Verificar">Verificar</span><span data-f="Confere">Confere</span><span data-f="todos">Todos</span>
<input id="qA" placeholder="Buscar produto, código ou NCM"></div>
<div style="max-height:680px;overflow:auto"><table><thead><tr><th>Situação</th><th>Produto</th><th>Como veio no SPED</th><th>O que a regra diz</th><th>Base legal</th><th class="r">Valor</th></tr></thead><tbody id="tA"></tbody></table></div>
<small id="nA" style="display:block;text-align:center;color:#64748b;margin-top:8px"></small><button class="mais" id="mA">Mostrar mais</button></div>
<div class="ft"><b>Tabelas usadas:</b> {_h.escape(fontes["st"])} · {_h.escape(fontes["aliq"])} · {_h.escape(fontes["pis"])}.<br><b>Ainda não incluído:</b> {_h.escape(fontes["pendente"])}.<br>
Análise de leitura: tudo é apontamento para o contador conferir. Nenhum dado do SPED é alterado.</div>
</section></div>
<script id="dados" type="application/json">{dados}</script><script>{_JS}</script></body></html>'''


# ------------------------------------------------------------------ Excel
def gerar_xlsx(res, caminho):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    cab_f, cab_p = Font(bold=True, color='FFFFFF'), PatternFill('solid', fgColor='1E3A8A')
    moeda, pct = '#,##0.00', '0.00"%"'
    e, t = res['empresa'], res['totais']

    def aba(nome, cab, linhas, larg, fmt):
        ws = wb.create_sheet(nome)
        ws.append(cab)
        for c in ws[1]:
            c.font, c.fill, c.alignment = cab_f, cab_p, Alignment(vertical='center', wrap_text=True)
        for l in linhas:
            ws.append(l)
        n = len(linhas) + 1
        if linhas:
            tot = ['Total'] + [''] * (len(cab) - 1)
            for j, f in enumerate(fmt):
                if f == moeda:
                    col = get_column_letter(j + 1)
                    tot[j] = f'=SUBTOTAL(9,{col}2:{col}{n})'
            ws.append(tot)
            for c in ws[n + 1]:
                c.font = Font(bold=True)
        for j, (w, f) in enumerate(zip(larg, fmt)):
            col = get_column_letter(j + 1)
            ws.column_dimensions[col].width = w
            if f:
                for r in range(2, n + 2):
                    ws[f'{col}{r}'].number_format = f
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = f'A1:{get_column_letter(len(cab))}{n}'
        return ws

    ws = wb.active
    ws.title = 'Resumo'
    d = lambda s: f'{s[:2]}/{s[2:4]}/{s[4:]}' if len(s or '') == 8 else (s or '')
    ws.append(['Tributação das entradas por produto · Análise Tributária (FiscoCont+)'])
    ws['A1'].font = Font(bold=True, size=14, color='1E3A8A')
    for l in [['Empresa', e.get('nome', '')], ['CNPJ', e.get('cnpj', '')], ['Período', f'{d(e.get("dt_ini", ""))} a {d(e.get("dt_fin", ""))}'], [],
              ['Notas de entrada', len(res['notas'])], ['Itens', len(res['itens'])], ['Produtos', len(res['produtos'])],
              ['Valor das entradas', t['vl']], ['Base do ICMS', t['bc']], ['ICMS creditado', t['icms']], ['ICMS-ST', t['st']], ['IPI', t['ipi']], [],
              ['Auditoria (produtos)', 'Quantidade', 'Valor']] + [[s, res['situacao'].get(s, [0, 0])[0], res['situacao'].get(s, [0, 0])[1]]
                                                                for s in ('Confere', 'Divergente', 'Verificar')] + [[], ['Pontos para conferir']] + \
             [[p[1]] for p in res['pontos']] + [[], ['Tabelas usadas'], [res['fontes']['st']], [res['fontes']['aliq']], [res['fontes']['pis']],
                                                ['Ainda não incluído: ' + res['fontes']['pendente']]]:
        ws.append(l)
    for r in range(9, 14):
        ws[f'B{r}'].number_format = moeda
    for r in range(16, 19):
        ws[f'C{r}'].number_format = moeda
    ws.column_dimensions['A'].width = 34
    ws.column_dimensions['B'].width = 22
    ws.column_dimensions['C'].width = 18
    N = res['notas']
    linhas = []
    for i in res['itens']:
        n = N[i['nota_idx']]
        linhas.append([n['num'], d(n['dt_es'] or n['dt']), n['forn'], n['cnpj'], n['uf'], n['chave'], i['cod'], i['desc'], i['ncm'],
                       TIPO_ITEM.get(i['tipo'], i['tipo']), i['cfop'], i['cst'], i['qtd'], i['unid'], i['vl'], i['bc'], i['aliq'], i['icms'],
                       i['bcst'], i['st'], i['ipi'], i['cstpis'], i['sit'],
                       ' | '.join(f'{a[1]}: {a[2]}' for a in i['ach'] if a[0] != 'Info'),
                       ' | '.join(a[2] for a in i['ach'] if a[0] == 'Info')])
    aba('Itens das notas', ['Nota', 'Entrada', 'Fornecedor', 'CNPJ fornecedor', 'UF', 'Chave', 'Código', 'Descrição', 'NCM', 'Tipo (0200)', 'CFOP',
                            'CST ICMS', 'Qtd', 'Un', 'Valor item', 'Base ICMS', 'Alíq. ICMS', 'ICMS', 'Base ST', 'ICMS-ST', 'IPI', 'CST PIS',
                            'Situação', 'Apontamentos', 'Informativos'], linhas,
        [10, 11, 32, 16, 5, 46, 16, 40, 11, 20, 7, 8, 10, 6, 13, 13, 9, 12, 12, 12, 11, 8, 11, 70, 50],
        [None] * 12 + ['#,##0.####', None, moeda, moeda, pct, moeda, moeda, moeda, moeda, None, None, None, None])
    aba('Por produto', ['Código', 'Descrição', 'NCM', 'Tipo (0200)', 'CFOPs', 'CSTs', 'Alíquotas', 'Notas', 'Qtd', 'Valor', 'Base ICMS', 'ICMS',
                        'ICMS-ST', 'IPI', 'Situação'],
        [[p['cod'], p['desc'], p['ncm'], TIPO_ITEM.get(p['tipo'], p['tipo']), '/'.join(p['cfops']), '/'.join(p['csts']),
          '/'.join(_pc(a) for a in p['aliqs']), p['notas'], p['qtd'], p['vl'], p['bc'], p['icms'], p['st'], p['ipi'], p['sit']] for p in res['produtos']],
        [16, 42, 11, 20, 14, 12, 14, 7, 10, 14, 14, 13, 13, 12, 11], [None] * 8 + ['#,##0.####', moeda, moeda, moeda, moeda, moeda, None])
    au = []
    for p in res['produtos']:
        for a in p['ach']:
            au.append([a[0] if a[0] != 'Info' else 'Informativo', a[1], p['cod'], p['desc'], p['ncm'], '/'.join(p['cfops']), '/'.join(p['csts']),
                       '/'.join(_pc(x) for x in p['aliqs']), a[2], a[3], p['vl']])
    au.sort(key=lambda x: ({'Divergente': 0, 'Verificar': 1}.get(x[0], 2), -x[10]))
    ws = aba('Auditoria', ['Situação', 'Tema', 'Código', 'Descrição', 'NCM', 'CFOPs', 'CSTs', 'Alíquotas', 'O que a regra diz', 'Base legal', 'Valor do produto'],
             au, [12, 13, 16, 40, 11, 12, 10, 12, 80, 40, 15], [None] * 10 + [moeda])
    cor = {'Divergente': 'FEE2E2', 'Verificar': 'FEF3C7', 'Informativo': 'E0E7FF'}
    for r in range(2, len(au) + 2):
        c = ws[f'A{r}']
        if c.value in cor:
            c.fill = PatternFill('solid', fgColor=cor[c.value])
    _aba_tipo_oper(wb, res, cab_f, cab_p, moeda)
    _aba_acum7000(wb, res, cab_f, cab_p, moeda)
    g = lambda lst, rot: [[rot(k), v['n'], v['vl'], v['bc'], v['icms'], v['st']] for k, v in lst]
    hd = ['Itens', 'Valor', 'Base ICMS', 'ICMS', 'ICMS-ST']
    fm5 = [None, None, moeda, moeda, moeda, moeda]
    aba('Por CFOP', ['CFOP'] + hd, g(res['por_cfop'], lambda c: f'{c} · {CFOP_NOME.get(c[1:], "")}'), [48, 8, 15, 15, 14, 14], fm5)
    aba('Por CST', ['CST'] + hd, g(res['por_cst'], lambda c: f'{c} · {CST_NOME.get((c or "")[-2:], "")}'), [42, 8, 15, 15, 14, 14], fm5)
    aba('Por alíquota', ['Alíquota'] + hd, g(res['por_aliq'], lambda a: _pc(a)), [14, 8, 15, 15, 14, 14], fm5)
    aba('Por fornecedor', ['Fornecedor'] + hd, g(res['por_forn'], lambda f: f), [46, 8, 15, 15, 14, 14], fm5)
    wb.save(caminho)


def _aba_tipo_oper(wb, res, cab_f, cab_p, moeda):
    """Aba "Por tipo de operação": resumo por tipo (revenda, uso e consumo, ativo...) e, abaixo, tipo → CFOP → produto
    com o agrupamento do Excel (botões 1/2/3)."""
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    ws = wb.create_sheet('Por tipo de operação')
    itens, notas = res['itens'], res['notas']
    grupos = OrderedDict((nome, []) for nome, _ in TIPO_OPER)
    grupos['Outras entradas'] = []
    for k, i in enumerate(itens):
        grupos[tipo_operacao(i['cfop'])].append(k)
    tot = sum(i['vl'] for i in itens) or 1
    ws.append(['Entradas por tipo de operação (pelo CFOP de cada item)'])
    ws['A1'].font = Font(bold=True, size=13, color='1E3A8A')
    ws.append(['Resumo no topo; abaixo, cada tipo abre por CFOP e por produto (botões 1/2/3 à esquerda).'])
    ws['A2'].font = Font(italic=True, color='475569')
    cab = ['Tipo de operação', 'CFOP', 'Código', 'Produto', 'NCM', 'Notas', 'Itens', 'Valor', '% das entradas', 'Base ICMS', 'ICMS', 'ICMS-ST', 'IPI',
           'Situação na auditoria']
    ws.append(cab)
    for c in ws[3]:
        c.font, c.fill, c.alignment = cab_f, cab_p, Alignment(vertical='center', wrap_text=True)
    for j, w in enumerate([34, 34, 16, 44, 11, 7, 7, 15, 10, 15, 13, 13, 12, 18]):
        ws.column_dimensions[get_column_letter(j + 1)].width = w
    soma = lambda ks, f: sum(itens[k][f] for k in ks)
    nn = lambda ks: len({itens[k]['nota_idx'] for k in ks})
    f_res, f_tipo, f_cf = PatternFill('solid', fgColor='EEF2FF'), PatternFill('solid', fgColor='DBEAFE'), PatternFill('solid', fgColor='F8FAFC')
    r = 3
    for nome, ks in grupos.items():           # resumo
        if not ks:
            continue
        r += 1
        ws.append([nome, ', '.join(sorted({itens[k]['cfop'] for k in ks})), '', '', '', nn(ks), len(ks), soma(ks, 'vl'), soma(ks, 'vl') / tot,
                   soma(ks, 'bc'), soma(ks, 'icms'), soma(ks, 'st'), soma(ks, 'ipi'),
                   f"{sum(1 for k in ks if itens[k]['sit'] == 'Divergente')} div. · {sum(1 for k in ks if itens[k]['sit'] == 'Verificar')} verif."])
        for c in ws[r]:
            c.fill = f_res
        ws.cell(r, 1).font = Font(bold=True)
    r += 1
    ws.append(['Total', '', '', '', '', len(notas), len(itens), tot if itens else 0, 1 if itens else 0, sum(i['bc'] for i in itens),
               sum(i['icms'] for i in itens), sum(i['st'] for i in itens), sum(i['ipi'] for i in itens), ''])
    for c in ws[r]:
        c.font = Font(bold=True)
    r += 1
    ws.append([])
    for nome, ks in grupos.items():           # detalhe
        if not ks:
            continue
        r += 1
        ws.append([nome, '', '', '', '', nn(ks), len(ks), soma(ks, 'vl'), soma(ks, 'vl') / tot, soma(ks, 'bc'), soma(ks, 'icms'), soma(ks, 'st'),
                   soma(ks, 'ipi'), ''])
        for c in ws[r]:
            c.fill, c.font = f_tipo, Font(bold=True, color='1E3A8A')
        por_cf = OrderedDict()
        for k in sorted(ks, key=lambda k: itens[k]['cfop']):
            por_cf.setdefault(itens[k]['cfop'], []).append(k)
        for cf, kc in por_cf.items():
            r += 1
            ws.append([nome, f"{cf} · {CFOP_NOME.get(cf[1:], '')}", '', '', '', nn(kc), len(kc), soma(kc, 'vl'), soma(kc, 'vl') / tot, soma(kc, 'bc'),
                       soma(kc, 'icms'), soma(kc, 'st'), soma(kc, 'ipi'), ''])
            ws.row_dimensions[r].outlineLevel = 1
            for c in ws[r]:
                c.fill = f_cf
            ws.cell(r, 2).font = Font(bold=True)
            ws.cell(r, 1).font = Font(color='94A3B8')
            por_p = OrderedDict()
            for k in kc:
                por_p.setdefault(itens[k]['cod'], []).append(k)
            for cod, kp in sorted(por_p.items(), key=lambda x: -soma(x[1], 'vl')):
                i0 = itens[kp[0]]
                sits = {itens[k]['sit'] for k in kp}
                sit = 'Divergente' if 'Divergente' in sits else 'Verificar' if 'Verificar' in sits else 'Confere'
                r += 1
                ws.append([nome, cf, cod, i0['desc'], i0['ncm'], nn(kp), len(kp), soma(kp, 'vl'), soma(kp, 'vl') / tot, soma(kp, 'bc'), soma(kp, 'icms'),
                           soma(kp, 'st'), soma(kp, 'ipi'), sit])
                ws.row_dimensions[r].outlineLevel = 2
                ws.cell(r, 1).font = Font(color='94A3B8')
                ws.cell(r, 2).font = Font(color='94A3B8')
                if sit != 'Confere':
                    ws.cell(r, 14).fill = PatternFill('solid', fgColor='FEE2E2' if sit == 'Divergente' else 'FEF3C7')
    for row in range(4, r + 1):
        for col in (8, 10, 11, 12, 13):
            ws.cell(row, col).number_format = moeda
        ws.cell(row, 9).number_format = '0.0%'
    ws.sheet_properties.outlinePr.summaryBelow = False
    ws.freeze_panes = 'A4'


def _aba_acum7000(wb, res, cab_f, cab_p, moeda):
    """Aba "Acumulador 7000 por NF": nota → acumulador sugerido → produtos (agrupamento do Excel em 3 níveis)."""
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    A = res.get('acum7000') or {'notas': []}
    ws = wb.create_sheet('Acumulador 7000 por NF')
    nomes = A.get('nomes', {})
    ns = A['notas']
    mist = sum(1 for n in ns if n['mista'])
    ws.append(['Acumulador 7000 (DESPESAS) por nota · sugestão do acumulador de despesa certo (7001 a 7040) pela descrição, NCM e CFOP'])
    ws['A1'].font = Font(bold=True, size=13, color='1E3A8A')
    ws.append([f'{len(ns)} nota(s) com itens no 7000 · R$ {_fm(sum(n["vl"] for n in ns))} · {len(ns) - mist} com um acumulador · {mist} mista(s) '
               '(produtos de naturezas diferentes: divida a nota ou escolha o acumulador — o sistema não escolhe). '
               'Botões 1/2/3 à esquerda: só notas · notas e acumuladores · tudo, com os produtos.'])
    ws['A2'].font = Font(italic=True, color='475569')
    cab = ['Nota', 'Entrada', 'Fornecedor', 'Valor no 7000', 'Situação / nível', 'Acumulador', 'Código do produto', 'Descrição',
           'NCM', 'CFOP', 'Motivo da classificação', 'Confiança', 'Valor', '% da NF']
    ws.append(cab)
    for c in ws[3]:
        c.font, c.fill, c.alignment = cab_f, cab_p, Alignment(vertical='center')
    larg = [11, 11, 34, 13, 26, 11, 18, 46, 11, 7, 40, 10, 13, 8]
    from openpyxl.utils import get_column_letter
    for j, w in enumerate(larg):
        ws.column_dimensions[get_column_letter(j + 1)].width = w
    f_mh, f_uh, f_ac = PatternFill('solid', fgColor='FDE6C4'), PatternFill('solid', fgColor='EEF2FF'), PatternFill('solid', fgColor='FFF7EC')
    cinza, azul, laranja = Font(color='94A3B8'), Font(bold=True, color='1E3A8A'), Font(bold=True, color='B45309')
    borda = Border(top=Side(style='medium', color='F59E0B'))
    d = lambda s: f'{s[:2]}/{s[2:4]}/{s[4:]}' if len(s or '') == 8 else (s or '')
    conf = {'alta': 'alta', 'media': 'média'}
    r = 3
    for n in ns:
        r += 1
        sit = f'⚠ Mista · {len(n["partes"])} acumuladores' if n['mista'] else 'Única'
        ws.append([n['num'], d(n['dt']), n['forn'], n['vl'], sit, '', '', 'divida a nota ou escolha o acumulador' if n['mista'] else '',
                   '', '', '', '', n['vl'], 1])
        for c in ws[r]:
            c.fill = f_mh if n['mista'] else f_uh
            c.font = Font(bold=True, color='B45309') if n['mista'] and c.column in (5, 8) else Font(bold=True)
            if n['mista']:
                c.border = borda
        for p in n['partes']:
            r += 1
            ws.append([n['num'], '', '', '', f'└ {p["acum"] or "sem sugestão"}', p['acum'], '', f'{nomes.get(p["acum"], "sem sugestão: classificar")} '
                       f'· {len(p["itens"])} produto(s)', '', '', '', conf.get(p['conf'], p['conf']), p['vl'], p['vl'] / n['vl'] if n['vl'] else 0])
            ws.row_dimensions[r].outlineLevel = 1
            for c in ws[r]:
                c.fill = f_ac
                c.font = azul if c.column in (5, 6, 8, 13, 14) else Font()
            ws.cell(r, 1).font = cinza
            for it in p['itens']:
                r += 1
                ws.append([n['num'], '', '', '', f'    · item {it["item"]}', p['acum'], it['cod'], it['desc'], it['ncm'], it['cfop'], it['motivo'],
                           conf.get(it['conf'], it['conf']), it['vl'], it['vl'] / n['vl'] if n['vl'] else 0])
                ws.row_dimensions[r].outlineLevel = 2
                ws.cell(r, 1).font = cinza
                ws.cell(r, 5).font = Font(color='64748B')
                ws.cell(r, 6).font = cinza
    for row in range(4, r + 1):
        ws.cell(row, 4).number_format = moeda
        ws.cell(row, 13).number_format = moeda
        ws.cell(row, 14).number_format = '0%'
    ws.sheet_properties.outlinePr.summaryBelow = False
    ws.freeze_panes = 'A4'
    ws.auto_filter.ref = f'A3:N{max(r, 4)}'
    if not ns:
        ws.append(['Nenhum item no acumulador 7000 neste SPED.'])


# ------------------------------------------------------------------ CLI
def main_cli(argv):
    def opt(n):
        return argv[argv.index(n) + 1] if n in argv and argv.index(n) + 1 < len(argv) else None
    sped, js = opt('--sped'), opt('--json')
    if not sped or not js:
        print('uso: aud-tributacao --sped ARQ.txt --json res.json [--html saida.html] [--xlsx saida.xlsx]')
        return 1
    out = {}
    try:
        import difal_ro as dr
        texto, _ = dr._ler(sped)
        if '|0110|' in texto[:20000] or '\n|M200|' in texto or '\n|M210|' in texto:
            raise ValueError('Esse arquivo é a EFD-Contribuições. Selecione o SPED Fiscal (EFD ICMS/IPI).')
        res = analisar_texto(texto)
        if not res['notas']:
            out = {'erro': 'Nenhuma nota de entrada (C100 com IND_OPER = 0) neste SPED.'}
        else:
            if opt('--html'):
                with open(opt('--html'), 'w', encoding='utf-8') as fh:
                    fh.write(gerar_html(res))
            if opt('--xlsx'):
                gerar_xlsx(res, opt('--xlsx'))
            e = res['empresa']
            out = {'ok': True, 'empresa': e.get('nome', ''), 'cnpj': e.get('cnpj', ''), 'periodo': [e.get('dt_ini', ''), e.get('dt_fin', '')],
                   'notas': len(res['notas']), 'itens': len(res['itens']), 'produtos': len(res['produtos']), 'situacao': res['situacao']}
    except ValueError as ex:
        out = {'erro': str(ex)}
    except Exception as ex:  # pragma: no cover
        import traceback
        out = {'erro': f'Falha ao processar: {ex}', 'detalhe': traceback.format_exc()[-1500:]}
    tmp = js + '.parcial'
    with open(tmp, 'w', encoding='utf-8') as fh:
        json.dump(out, fh, ensure_ascii=False)
    os.replace(tmp, js)
    return 0


if __name__ == '__main__':
    sys.exit(main_cli(sys.argv[1:]))
